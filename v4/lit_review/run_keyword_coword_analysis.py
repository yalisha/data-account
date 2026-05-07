from __future__ import annotations

import csv
import math
import re
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path("/Users/mac/computerscience/0做完了/15会计研究/v4/lit_review")
POOL_XLSX = BASE / "个人破产制度下载文献池_初筛_v1.xlsx"
OUT_DIR = BASE / "keyword_coword_20260506"
OUT_XLSX = OUT_DIR / "个人破产制度关键词共联分析_v1.xlsx"
OUT_REPORT = OUT_DIR / "个人破产制度关键词共联分析_结论.md"
OUT_SVG = OUT_DIR / "个人破产制度概念共现网络.svg"
OUT_GAPS = OUT_DIR / "个人破产制度潜在空白组合.csv"


CONCEPTS: dict[str, list[str]] = {
    "PB制度": [
        "personal bankruptcy",
        "bankruptcy exemption",
        "personal bankruptcy law",
        "debtor protection",
        "fresh start policy",
        "个人破产",
        "个人债务集中清理",
        "类个人破产",
        "破产免责",
    ],
    "自然人保证/连带责任": [
        "personal guarantee",
        "shareholder guarantee",
        "natural person guarantee",
        "guarantor",
        "joint liability",
        "personal liability",
        "保证责任",
        "连带责任保证",
        "个人保证",
        "自然人担保",
        "经营者保证",
        "企业经营者保证",
        "实际控制人担保",
        "控股股东担保",
        "法定代表人担保",
        "董监高担保",
        "连带责任",
    ],
    "控制人失败成本": [
        "controlling shareholder",
        "ultimate controller",
        "founder",
        "family firm",
        "private entrepreneur",
        "entrepreneurial failure",
        "failure cost",
        "失败成本",
        "实际控制人",
        "自然人实控",
        "创始人",
        "家族企业",
        "民营企业",
        "企业家",
        "控制人",
    ],
    "股权质押/控制权风险": [
        "share pledge",
        "share pledging",
        "equity pledge",
        "pledged shares",
        "control rights",
        "股权质押",
        "股份质押",
        "控制权风险",
        "控制权",
    ],
    "银行信贷供给": [
        "credit supply",
        "bank lending",
        "loan supply",
        "credit rationing",
        "credit access",
        "信贷资源供给",
        "银行信贷",
        "贷款供给",
        "信贷投放",
        "信贷可得",
    ],
    "债务期限/融资结构": [
        "debt maturity",
        "short-term debt",
        "long-term debt",
        "debt structure",
        "capital structure",
        "financing structure",
        "debt financing",
        "债务期限",
        "债务结构",
        "融资结构",
        "资本结构",
        "短期贷款",
        "长期贷款",
        "短期借款",
        "长期借款",
        "短贷长投",
    ],
    "坏消息隐藏/崩盘风险": [
        "stock price crash",
        "crash risk",
        "bad news hoarding",
        "bad news",
        "information risk",
        "坏消息隐藏",
        "股价崩盘",
        "崩盘风险",
        "信息风险",
        "风险累积",
    ],
    "风险披露/文本具体性": [
        "risk disclosure",
        "disclosure specificity",
        "specificity",
        "textual disclosure",
        "annual report risk",
        "boilerplate",
        "template",
        "信息披露",
        "风险披露",
        "披露质量",
        "文本披露",
        "具体性",
        "模板化",
        "文本相似",
    ],
    "会计稳健/损失确认": [
        "accounting conservatism",
        "conditional conservatism",
        "bad news timeliness",
        "loss recognition",
        "timely loss",
        "会计稳健",
        "条件稳健",
        "坏消息确认",
        "损失确认",
        "及时确认",
    ],
    "资产减值/信用减值": [
        "asset impairment",
        "goodwill impairment",
        "credit impairment",
        "impairment loss",
        "loss allowance",
        "资产减值",
        "信用减值",
        "商誉减值",
        "减值准备",
        "坏账准备",
        "预期信用损失",
    ],
    "盈余管理/盈余平滑": [
        "earnings management",
        "income smoothing",
        "real earnings management",
        "accruals",
        "盈余管理",
        "盈余平滑",
        "真实盈余管理",
        "应计",
        "操纵",
    ],
    "审计/KAM": [
        "audit",
        "audit fee",
        "audit fees",
        "auditor",
        "key audit matters",
        "going concern",
        "审计",
        "审计费用",
        "审计意见",
        "关键审计事项",
        "持续经营",
        "非标意见",
    ],
    "创新/研发": [
        "innovation",
        "patent",
        "r&d",
        "research and development",
        "科技创业",
        "创新",
        "专利",
        "研发",
        "新质生产力",
    ],
    "创业/fresh-start": [
        "entrepreneurship",
        "entrepreneurial",
        "startup",
        "startups",
        "self-employment",
        "fresh start",
        "创业",
        "再创业",
        "自雇",
        "企业家精神",
    ],
    "家庭/居民金融": [
        "household",
        "consumer",
        "consumption",
        "household debt",
        "居民",
        "家庭",
        "消费",
        "预防性储蓄",
        "家庭债务",
    ],
    "法学制度/司法实践": [
        "legislation",
        "court",
        "judicial",
        "law",
        "bankruptcy court",
        "legal",
        "立法",
        "法院",
        "司法",
        "执行",
        "制度构建",
        "破产法庭",
        "税收",
        "法律",
    ],
    "员工/劳动": [
        "employee",
        "labor",
        "labour",
        "union",
        "wage",
        "员工",
        "职工",
        "劳动",
        "工会",
        "薪酬",
    ],
    "商业信用/供应链": [
        "trade credit",
        "supplier",
        "customer",
        "commercial credit",
        "supply chain",
        "商业信用",
        "供应商",
        "客户",
        "供应链",
        "应付账款",
        "应付票据",
    ],
    "风险承担/投资": [
        "risk taking",
        "investment",
        "overinvestment",
        "underinvestment",
        "investment efficiency",
        "风险承担",
        "投资",
        "过度投资",
        "投资不足",
        "投资效率",
    ],
}

EXPOSURE_CONCEPTS = [
    "自然人保证/连带责任",
    "控制人失败成本",
    "股权质押/控制权风险",
    "银行信贷供给",
    "债务期限/融资结构",
]

Y_CONCEPTS = [
    "坏消息隐藏/崩盘风险",
    "风险披露/文本具体性",
    "会计稳健/损失确认",
    "资产减值/信用减值",
    "盈余管理/盈余平滑",
    "审计/KAM",
    "创新/研发",
    "风险承担/投资",
    "商业信用/供应链",
    "债务期限/融资结构",
]

TOPIC_CANDIDATES = [
    (
        "自然人保证/连带责任",
        "风险披露/文本具体性",
        "自然人担保暴露下，PB 是否提高风险披露具体性",
        "主线候选：最贴近坏消息隐藏过程，若 count 很低就是优先空白。",
    ),
    (
        "自然人保证/连带责任",
        "会计稳健/损失确认",
        "自然人担保暴露下，PB 是否提高坏消息确认及时性/会计稳健性",
        "会计味强，适合替代 crash risk 做主 Y。",
    ),
    (
        "自然人保证/连带责任",
        "资产减值/信用减值",
        "自然人担保暴露下，PB 是否促进减值及时确认",
        "更具体，但受行业和疫情冲击影响大。",
    ),
    (
        "自然人保证/连带责任",
        "审计/KAM",
        "自然人追偿价值下降后，审计师是否增加债务/KAM 风险披露",
        "可做机制或审计方向副主线。",
    ),
    (
        "控制人失败成本",
        "风险披露/文本具体性",
        "控制人失败成本下降后，企业风险披露是否更具体",
        "坏消息隐藏上游 Y，适合会计/管理交叉。",
    ),
    (
        "控制人失败成本",
        "会计稳健/损失确认",
        "控制人失败成本下降后，企业是否更及时确认坏消息",
        "会计主线候选。",
    ),
    (
        "股权质押/控制权风险",
        "坏消息隐藏/崩盘风险",
        "PB 是否缓解质押控制人的坏消息隐藏和崩盘风险",
        "相邻 crash/pledge 文献较多，需谨慎。",
    ),
    (
        "债务期限/融资结构",
        "盈余管理/盈余平滑",
        "债务短期化压力是否诱发现实经营或盈余管理",
        "可能被信贷和盈余管理文献挤压，作为机制更好。",
    ),
    (
        "银行信贷供给",
        "商业信用/供应链",
        "PB 后银行收缩是否引发商业信用替代",
        "金融机制候选，离会计主线较远。",
    ),
    (
        "控制人失败成本",
        "创新/研发",
        "失败成本下降是否提高存量企业创新/研发",
        "已接近创业创新文献，要做必须写双机制。",
    ),
]


def normalize_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def load_pool_records() -> list[dict]:
    wb = load_workbook(POOL_XLSX, read_only=True, data_only=True)
    ws = wb["全部去重池"]
    headers = [normalize_text(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = {headers[i]: row[i] for i in range(len(headers))}
        try:
            rec["相关度分"] = float(rec.get("相关度分") or 0)
        except Exception:
            rec["相关度分"] = 0.0
        records.append(rec)
    return records


def tag_record(rec: dict) -> set[str]:
    blob = " ".join(
        normalize_text(rec.get(k))
        for k in ["标题", "关键词", "摘要", "文献类别", "与本题关系", "撞题判断"]
    ).lower()
    tags = set()
    for concept, terms in CONCEPTS.items():
        for term in terms:
            if term.lower() in blob:
                tags.add(concept)
                break
    return tags


def relevant_records(records: list[dict]) -> list[dict]:
    out = []
    for rec in records:
        category = normalize_text(rec.get("文献类别"))
        if category != "低相关/待剔除" or rec.get("相关度分", 0) >= 8:
            out.append(rec)
    return out


def strict_pb_records(records: list[dict], tags_by_idx: dict[int, set[str]]) -> list[dict]:
    out = []
    for idx, rec in enumerate(records):
        if "PB制度" in tags_by_idx[idx] and normalize_text(rec.get("文献类别")) != "低相关/待剔除":
            out.append(rec)
    return out


def concept_stats(records: list[dict], tags_by_key: dict[int, set[str]], index_map: dict[int, int]) -> tuple[Counter, Counter]:
    counts: Counter = Counter()
    pairs: Counter = Counter()
    for local_idx, rec in enumerate(records):
        global_idx = index_map[local_idx]
        tags = sorted(tags_by_key[global_idx])
        for tag in tags:
            counts[tag] += 1
        for a, b in combinations(tags, 2):
            pairs[(a, b)] += 1
    return counts, pairs


def pair_count(pairs: Counter, a: str, b: str) -> int:
    if a == b:
        return 0
    key = tuple(sorted([a, b]))
    return pairs.get(key, 0)


def make_matrix(pairs: Counter, rows: list[str], cols: list[str]) -> list[list[object]]:
    matrix = [[""] + cols]
    for row_concept in rows:
        matrix.append([row_concept] + [pair_count(pairs, row_concept, col) for col in cols])
    return matrix


def write_table(ws, rows: list[list[object]], freeze: str = "A2") -> None:
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = freeze
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 20 if idx > 1 else 28


def candidate_rows(strict_pairs: Counter, broad_pairs: Counter, strict_counts: Counter, broad_counts: Counter) -> list[list[object]]:
    rows = [
        [
            "候选组合",
            "Exposure概念",
            "Y/机制概念",
            "严格PB池共现数",
            "全相关池共现数",
            "PB+Y共现数",
            "Exposure+Y共现数",
            "空白判断",
            "备注",
        ]
    ]
    for exp, y, title, note in TOPIC_CANDIDATES:
        strict_xy = pair_count(strict_pairs, exp, y)
        broad_xy = pair_count(broad_pairs, exp, y)
        pb_y = pair_count(strict_pairs, "PB制度", y)
        exp_y = pair_count(broad_pairs, exp, y)
        if strict_xy == 0:
            verdict = "本批严格PB池未见共现"
        elif strict_xy <= 2:
            verdict = "极少共现，可人工核查"
        elif strict_xy <= 5:
            verdict = "少量共现，需精查"
        else:
            verdict = "已有较多共现，不优先当空白"
        rows.append([title, exp, y, strict_xy, broad_xy, pb_y, exp_y, verdict, note])
    return rows


def top_record_rows(records: list[dict], tags_by_idx: dict[int, set[str]], index_map: dict[int, int], concept_a: str, concept_b: str, limit: int = 15) -> list[list[object]]:
    rows = [["概念A", "概念B", "年份", "标题", "来源", "文献类别", "相关度分"]]
    hits = []
    for local_idx, rec in enumerate(records):
        global_idx = index_map[local_idx]
        tags = tags_by_idx[global_idx]
        if concept_a in tags and concept_b in tags:
            hits.append(rec)
    hits = sorted(hits, key=lambda r: (-float(r.get("相关度分") or 0), normalize_text(r.get("年份")), normalize_text(r.get("标题"))))
    for rec in hits[:limit]:
        rows.append(
            [
                concept_a,
                concept_b,
                rec.get("年份"),
                rec.get("标题"),
                rec.get("刊物/来源"),
                rec.get("文献类别"),
                rec.get("相关度分"),
            ]
        )
    return rows


def export_svg(counts: Counter, pairs: Counter) -> None:
    concepts = [c for c, n in counts.most_common() if n >= 4]
    if not concepts:
        return
    width, height = 1200, 900
    cx, cy = width / 2, height / 2
    radius = 330
    positions = {}
    for i, concept in enumerate(concepts):
        angle = 2 * math.pi * i / len(concepts) - math.pi / 2
        positions[concept] = (cx + radius * math.cos(angle), cy + radius * math.sin(angle))

    max_count = max(counts[c] for c in concepts)
    max_edge = max([v for (a, b), v in pairs.items() if a in concepts and b in concepts] or [1])

    lines = [
        '<svg xmlns="http://www.w3.org/2000/svg" width="1200" height="900" viewBox="0 0 1200 900">',
        '<rect width="1200" height="900" fill="#ffffff"/>',
        '<text x="40" y="42" font-size="24" font-family="Arial" font-weight="bold">个人破产制度概念共现网络</text>',
        '<text x="40" y="72" font-size="14" font-family="Arial" fill="#555">节点大小=概念出现次数；边宽=共现次数；仅显示出现次数>=4的概念</text>',
    ]
    for (a, b), weight in pairs.items():
        if a not in positions or b not in positions or weight < 2:
            continue
        x1, y1 = positions[a]
        x2, y2 = positions[b]
        stroke = 0.6 + 5.0 * weight / max_edge
        opacity = min(0.85, 0.18 + weight / max_edge)
        lines.append(
            f'<line x1="{x1:.1f}" y1="{y1:.1f}" x2="{x2:.1f}" y2="{y2:.1f}" '
            f'stroke="#8aa6bf" stroke-width="{stroke:.2f}" opacity="{opacity:.2f}"/>'
        )
    for concept in concepts:
        x, y = positions[concept]
        size = 14 + 30 * math.sqrt(counts[concept] / max_count)
        color = "#C00000" if concept in ["PB制度", "自然人保证/连带责任", "控制人失败成本"] else "#1F77B4"
        lines.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="{size:.1f}" fill="{color}" opacity="0.82"/>')
        label = concept
        text_x = x + size + 5 if x < cx else x - size - 5
        anchor = "start" if x < cx else "end"
        lines.append(
            f'<text x="{text_x:.1f}" y="{y+4:.1f}" text-anchor="{anchor}" '
            f'font-size="15" font-family="Arial" fill="#111">{label} ({counts[concept]})</text>'
        )
    lines.append("</svg>")
    OUT_SVG.write_text("\n".join(lines), encoding="utf-8")


def write_report(
    total_records: int,
    broad_records: list[dict],
    strict_records: list[dict],
    strict_counts: Counter,
    strict_pairs: Counter,
    broad_pairs: Counter,
    candidates: list[list[object]],
) -> None:
    high_potential = candidates[1:]
    high_potential = sorted(high_potential, key=lambda r: (int(r[3]), int(r[4]), str(r[0])))

    lines = [
        "# 个人破产制度关键词共联分析",
        "",
        "更新日期：2026-05-06",
        "",
        "## 一、口径说明",
        "",
        f"- 原始去重文献池：{total_records} 条。",
        f"- 全相关池：{len(broad_records)} 条，剔除上一轮标记为低相关/待剔除的噪声文献。",
        f"- 严格 PB 池：{len(strict_records)} 条，要求同时被标注为相关且文本中出现个人破产、个人债务集中清理、类个人破产、personal bankruptcy、bankruptcy exemption 等核心概念。",
        "",
        "注意：这里的“没人做过”只表示在主人当前下载的 Scopus/CNKI 元数据池中未见概念共现，不等于全网绝对没有文章。它适合用来确定下一轮精查和实证优先级。",
        "",
        "## 二、严格 PB 池概念频率",
        "",
    ]
    for concept, count in strict_counts.most_common(20):
        lines.append(f"- {concept}：{count}")

    lines += [
        "",
        "## 三、最值得精查的低共现组合",
        "",
    ]
    for row in high_potential[:8]:
        lines.append(f"- {row[0]}：严格PB池共现 {row[3]}，全相关池共现 {row[4]}。判断：{row[7]}。")

    lines += [
        "",
        "## 四、初步结论",
        "",
        "从共联结果看，已经相对拥挤的是：个人破产与信贷/债务、个人破产与创业创新、个人破产与家庭/居民金融、个人破产与法学制度。它们更适合作背景或机制，不适合作新的主 Y。",
        "",
        "相对空白且更适合会计主线的是：自然人保证/连带责任与风险披露具体性、会计稳健/损失确认、资产减值/信用减值、审计/KAM 的交叉。这些组合更贴近“坏消息隐藏过程”，比直接拿 NCSKEW/DUVOL 当唯一主 Y 更能避开已有文献。",
        "",
        "因此，下一轮不应继续泛跑新 Y，而应优先补自然人担保或股权质押数据，并把主 Y 从 crash risk 前移到风险披露、稳健性和减值及时性。",
        "",
        "## 五、产出文件",
        "",
        f"- 共联工作簿：{OUT_XLSX}",
        f"- 空白组合 CSV：{OUT_GAPS}",
        f"- 共现网络 SVG：{OUT_SVG}",
    ]
    OUT_REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    all_records = load_pool_records()
    tags_by_idx = {idx: tag_record(rec) for idx, rec in enumerate(all_records)}

    broad = relevant_records(all_records)
    broad_indices = [all_records.index(rec) for rec in broad]
    broad_index_map = {i: global_idx for i, global_idx in enumerate(broad_indices)}

    strict = strict_pb_records(all_records, tags_by_idx)
    strict_indices = [all_records.index(rec) for rec in strict]
    strict_index_map = {i: global_idx for i, global_idx in enumerate(strict_indices)}

    broad_counts, broad_pairs = concept_stats(broad, tags_by_idx, broad_index_map)
    strict_counts, strict_pairs = concept_stats(strict, tags_by_idx, strict_index_map)

    candidates = candidate_rows(strict_pairs, broad_pairs, strict_counts, broad_counts)

    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("概念频率_严格PB池")
    write_table(ws, [["概念", "出现文献数"]] + [[k, v] for k, v in strict_counts.most_common()])

    ws = wb.create_sheet("概念频率_全相关池")
    write_table(ws, [["概念", "出现文献数"]] + [[k, v] for k, v in broad_counts.most_common()])

    ws = wb.create_sheet("共现边_严格PB池")
    strict_edge_rows = [["概念A", "概念B", "共现数"]]
    for (a, b), n in strict_pairs.most_common():
        strict_edge_rows.append([a, b, n])
    write_table(ws, strict_edge_rows)

    ws = wb.create_sheet("Exposure_Y矩阵_严格PB池")
    write_table(ws, make_matrix(strict_pairs, EXPOSURE_CONCEPTS, Y_CONCEPTS))

    ws = wb.create_sheet("Exposure_Y矩阵_全相关池")
    write_table(ws, make_matrix(broad_pairs, EXPOSURE_CONCEPTS, Y_CONCEPTS))

    ws = wb.create_sheet("潜在空白组合")
    write_table(ws, candidates)

    key_pairs = [
        ("自然人保证/连带责任", "风险披露/文本具体性"),
        ("自然人保证/连带责任", "会计稳健/损失确认"),
        ("自然人保证/连带责任", "审计/KAM"),
        ("控制人失败成本", "风险披露/文本具体性"),
        ("控制人失败成本", "会计稳健/损失确认"),
        ("PB制度", "坏消息隐藏/崩盘风险"),
        ("PB制度", "盈余管理/盈余平滑"),
        ("PB制度", "债务期限/融资结构"),
        ("PB制度", "创新/研发"),
    ]
    ws = wb.create_sheet("关键组合命中文献")
    rows = [["概念A", "概念B", "年份", "标题", "来源", "文献类别", "相关度分"]]
    for a, b in key_pairs:
        rows.extend(top_record_rows(strict, tags_by_idx, strict_index_map, a, b, limit=8)[1:])
    write_table(ws, rows)

    wb.save(OUT_XLSX)

    with OUT_GAPS.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(candidates)

    export_svg(strict_counts, strict_pairs)
    write_report(len(all_records), broad, strict, strict_counts, strict_pairs, broad_pairs, candidates)

    print(f"total={len(all_records)} broad={len(broad)} strict={len(strict)}")
    print(OUT_XLSX)
    print(OUT_REPORT)
    print(OUT_SVG)
    print(OUT_GAPS)


if __name__ == "__main__":
    main()
