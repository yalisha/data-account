from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


BASE = Path("/Users/mac/computerscience/0做完了/15会计研究")
RAW_DIR = BASE / "bib/v4个人破产制度"
OUT_DIR = BASE / "v4/lit_review"
POOL_XLSX = OUT_DIR / "个人破产制度下载文献池_初筛_v1.xlsx"
REVIEW_V1 = OUT_DIR / "个人破产制度文献综述_v1.xlsx"
REVIEW_V2 = OUT_DIR / "个人破产制度文献综述_v2.xlsx"
REPORT = OUT_DIR / "个人破产制度下载文献初筛报告.md"


REVIEW_HEADERS = [
    "年份",
    "APA格式的引用",
    "文献类别",
    "X",
    "自然人责任/暴露测度",
    "Y",
    "M",
    "模型",
    "研究问题",
    "摘要",
    "与本题关系",
    "撞题判断",
    "下一步用途/备注",
]


POOL_HEADERS = [
    "来源库",
    "原始文件",
    "年份",
    "标题",
    "作者",
    "刊物/来源",
    "DOI/链接",
    "关键词",
    "摘要",
    "文献类别",
    "相关度分",
    "与本题关系",
    "撞题判断",
    "建议动作",
]


def norm_text(value: str | None) -> str:
    return re.sub(r"\s+", " ", (value or "").strip())


def lower_blob(record: dict) -> str:
    return " ".join(str(v or "") for v in record.values()).lower()


def contains_any(text: str, terms: list[str]) -> bool:
    lowered = text.lower()
    return any(term.lower() in lowered for term in terms)


def parse_scopus_csv(path: Path) -> list[dict]:
    rows: list[dict] = []
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(
                {
                    "source_db": "Scopus",
                    "raw_file": path.name,
                    "year": norm_text(row.get("Year")),
                    "title": norm_text(row.get("Title")),
                    "authors": norm_text(row.get("Authors")),
                    "source": norm_text(row.get("Source title")),
                    "doi_link": norm_text(row.get("DOI")) or norm_text(row.get("Link")),
                    "keywords": "; ".join(
                        x
                        for x in [
                            norm_text(row.get("Author Keywords")),
                            norm_text(row.get("Index Keywords")),
                        ]
                        if x
                    ),
                    "abstract": norm_text(row.get("Abstract")),
                    "cited_by": norm_text(row.get("Cited by")),
                    "raw": row,
                }
            )
    return rows


def parse_cnki_txt(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="replace")
    chunks = re.split(r"(?=SrcDatabase-来源库:)", text)
    field_map = {
        "Title-题名": "title",
        "Author-作者": "authors",
        "Source-文献来源": "source",
        "Keyword-关键词": "keywords",
        "Summary-摘要": "abstract",
        "PubTime-发表时间": "pubtime",
        "DOI-DOI": "doi_link",
    }
    records: list[dict] = []
    for chunk in chunks:
        if "Title-题名:" not in chunk:
            continue
        rec = {
            "source_db": "CNKI",
            "raw_file": path.name,
            "year": "",
            "title": "",
            "authors": "",
            "source": "",
            "doi_link": "",
            "keywords": "",
            "abstract": "",
            "cited_by": "",
            "raw": {},
        }
        current_key = None
        for line in chunk.splitlines():
            if "-" in line and ":" in line:
                prefix, value = line.split(":", 1)
                prefix = prefix.strip()
                value = value.strip()
                if prefix in field_map:
                    current_key = field_map[prefix]
                    rec[current_key] = norm_text(value)
                else:
                    current_key = None
            elif current_key and line.strip():
                rec[current_key] = norm_text(rec[current_key] + " " + line.strip())
        if rec.get("pubtime"):
            match = re.search(r"(19|20)\d{2}", rec["pubtime"])
            if match:
                rec["year"] = match.group(0)
        if not rec["year"]:
            match = re.search(r"(19|20)\d{2}", chunk)
            if match:
                rec["year"] = match.group(0)
        records.append(rec)
    return records


def classify(record: dict) -> tuple[str, int, str, str, str]:
    title = record["title"]
    source = record["source"]
    keywords = record["keywords"]
    abstract = record["abstract"]
    blob = f"{title} {source} {keywords} {abstract}".lower()
    cn_blob = f"{title} {source} {keywords} {abstract}"

    score = 0
    category = "低相关/待剔除"
    relation = "暂未显示与企业会计/融资/创新直接相关。"
    collision = "低相关。"
    action = "除非后续综述需要制度背景，否则不优先读。"

    pb_en = contains_any(blob, ["personal bankruptcy", "bankruptcy exemption", "bankruptcy law"])
    fresh_start = "fresh start" in blob
    pb_cn = contains_any(cn_blob, ["个人破产", "个人债务集中清理", "类个人破产"])

    corp = contains_any(
        blob,
        [
            "corporate",
            "firm",
            "company",
            "entrepreneur",
            "startup",
            "capital structure",
            "investment",
            "innovation",
            "labor",
            "bank",
            "credit",
            "loan",
            "debt",
            "audit",
            "accounting",
            "disclosure",
            "earnings",
        ],
    ) or contains_any(cn_blob, ["企业", "公司", "上市公司", "银行", "信贷", "融资", "创业", "创新", "审计", "会计"])

    if pb_en:
        score += 5
    if pb_cn:
        score += 6
    if fresh_start:
        score += 1
    if corp:
        score += 3

    if contains_any(blob, ["earnings management", "income smoothing", "accounting conservatism", "audit fee", "audit fees"]) or contains_any(cn_blob, ["盈余管理", "盈余平滑", "会计稳健", "审计费用", "审计意见", "关键审计事项"]):
        score += 5
    if contains_any(blob, ["stock price crash", "crash risk", "bad news hoarding"]) or contains_any(cn_blob, ["股价崩盘", "坏消息隐藏"]):
        score += 5
    if contains_any(blob, ["credit supply", "bank lending", "loan supply", "debt maturity", "collateral", "guarantee"]) or contains_any(cn_blob, ["信贷资源供给", "债务期限", "债务结构", "抵押", "质押", "保证贷款", "担保"]):
        score += 4
    if contains_any(blob, ["entrepreneurship", "entrepreneurial", "risk taking", "innovation", "startup"]) or contains_any(cn_blob, ["创业", "创新", "风险承担", "企业家"]):
        score += 4
    if contains_any(blob, ["personal guarantee", "shareholder guarantee", "controlling shareholder", "share pledge", "personal liability"]) or contains_any(cn_blob, ["自然人担保", "实际控制人", "控股股东", "股权质押", "个人保证", "连带责任"]):
        score += 5

    unrelated_fresh = fresh_start and not pb_en and not contains_any(blob, ["bankruptcy", "credit", "debt", "entrepreneur", "firm", "corporate", "company"])
    if unrelated_fresh:
        score -= 5

    if contains_any(title, ["个人破产制度与企业盈余平滑"]):
        category = "中国PB-企业会计/资本市场"
        relation = "直接覆盖中国PB、企业盈余平滑，并把股价崩盘风险作为经济后果。"
        collision = "强撞题；不能复刻PB->盈余平滑或裸PB->崩盘风险。"
        action = "全文精读，作为边界文献和必须区分对象。"
        score += 10
    elif contains_any(title, ["个人破产制度与信贷资源供给"]):
        category = "中国PB-银行信贷"
        relation = "直接覆盖PB对银行信贷供给、期限和担保结构的影响。"
        collision = "强相邻；B线不能停在信贷收缩。"
        action = "全文精读，支撑信贷收缩竞争机制。"
        score += 10
    elif contains_any(title, ["个人破产试点促进城市科技创业活跃度"]):
        category = "中国PB-创业创新"
        relation = "直接覆盖PB试点与城市科技创业活跃度。"
        collision = "相邻；C线不能写城市创业数量。"
        action = "全文精读，支撑失败容错背景。"
        score += 8
    elif pb_cn and contains_any(cn_blob, ["居民消费", "家庭", "预防性储蓄", "家庭金融"]):
        category = "中国PB-家庭金融"
        relation = "提供中国PB影响居民风险偏好、消费和家庭债务的制度后果。"
        collision = "相邻但不撞企业会计。"
        action = "可读摘要，选择性引用制度背景。"
    elif pb_cn and contains_any(cn_blob, ["银行", "信贷", "融资", "债务", "担保", "抵押", "质押"]):
        category = "中国PB-银行/融资"
        relation = "支撑PB改变追偿价值、信贷条件和债务结构。"
        collision = "相邻；看是否企业层面。"
        action = "优先筛是否有企业微观结果。"
    elif pb_cn and contains_any(cn_blob, ["审计", "会计", "盈余", "股价崩盘", "信息披露"]):
        category = "中国PB-会计/资本市场"
        relation = "可能直接进入本题企业会计Y。"
        collision = "需人工核验是否撞题。"
        action = "优先下载全文。"
    elif pb_cn and contains_any(cn_blob, ["创业", "创新", "企业家"]):
        category = "中国PB-创业创新"
        relation = "支撑失败容错和企业家风险承担机制。"
        collision = "相邻；看层级是城市还是企业。"
        action = "优先读摘要和模型。"
    elif pb_cn:
        category = "中国PB-法学制度"
        relation = "制度背景、政策口径和司法实践材料。"
        collision = "通常不撞实证题。"
        action = "留作制度背景，不优先精读。"
    elif pb_en and contains_any(blob, ["corporate policies", "capital structure", "labor", "investment"]):
        category = "英文PB-公司政策"
        relation = "证明个人破产制度可以通过自然人主体影响公司政策。"
        collision = "理论相邻，不撞中国企业会计。"
        action = "优先精读。"
    elif pb_en and contains_any(blob, ["credit", "loan", "bank", "debt", "collateral"]):
        category = "英文PB-信贷/债务契约"
        relation = "支撑债权追偿价值和信贷重新定价机制。"
        collision = "机制基础，不撞。"
        action = "优先纳入综述。"
    elif (pb_en or fresh_start) and contains_any(blob, ["entrepreneur", "startup", "innovation", "risk taking"]):
        category = "英文PB-创业/失败成本"
        relation = "支撑失败保险、风险承担和创新机制。"
        collision = "机制基础，不撞。"
        action = "优先纳入综述。"
    elif contains_any(blob, ["personal guarantee", "shareholder guarantee", "controlling shareholder", "share pledge", "crash risk", "bad news hoarding"]) or contains_any(cn_blob, ["自然人担保", "控股股东", "股权质押", "坏消息隐藏", "股价崩盘"]):
        category = "自然人责任/坏消息隐藏相邻"
        relation = "可服务Exposure或Y的理论构造。"
        collision = "相邻理论，不撞PB主线。"
        action = "按是否能支持自然人责任边界筛选。"

    if score >= 13 and collision == "低相关。":
        collision = "高相关，需人工判断是否撞X/Y/机制。"
        action = "优先读摘要，必要时下载全文。"
    elif score >= 8 and collision == "低相关。":
        collision = "相邻或背景文献。"
        action = "保留在文献池，二轮筛。"

    return category, max(score, 0), relation, collision, action


def dedupe(records: list[dict]) -> list[dict]:
    seen: set[str] = set()
    unique: list[dict] = []
    for rec in records:
        doi = rec.get("doi_link", "").strip().lower()
        title = re.sub(r"\W+", "", rec.get("title", "").lower())
        key = doi if doi else title
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(rec)
    return unique


def enrich(records: list[dict]) -> list[dict]:
    out: list[dict] = []
    for rec in records:
        category, score, relation, collision, action = classify(rec)
        rec = dict(rec)
        rec.update(
            {
                "category": category,
                "score": score,
                "relation": relation,
                "collision": collision,
                "action": action,
            }
        )
        out.append(rec)
    return sorted(out, key=lambda x: (-x["score"], x.get("year", ""), x.get("title", "")))


def write_rows(ws, headers: list[str], rows: list[list[str]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        1: 12,
        2: 28,
        3: 10,
        4: 42,
        5: 24,
        6: 24,
        7: 26,
        8: 32,
        9: 60,
        10: 22,
        11: 10,
        12: 42,
        13: 36,
        14: 34,
    }
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(idx, 22)


def pool_row(rec: dict) -> list[str]:
    return [
        rec["source_db"],
        rec["raw_file"],
        rec["year"],
        rec["title"],
        rec["authors"],
        rec["source"],
        rec["doi_link"],
        rec["keywords"],
        rec["abstract"],
        rec["category"],
        rec["score"],
        rec["relation"],
        rec["collision"],
        rec["action"],
    ]


def apa_like(rec: dict) -> str:
    authors = rec["authors"] or "Unknown author"
    year = rec["year"] or "n.d."
    title = rec["title"]
    source = rec["source"]
    doi = rec["doi_link"]
    tail = f". {doi}" if doi else ""
    return f"{authors} ({year}). {title}. {source}{tail}"


def review_row_from_pool(rec: dict) -> list[str]:
    category = rec["category"]
    blob = lower_blob(rec)
    x = "个人破产制度/个人债务集中清理/破产豁免或fresh start制度。"
    exposure = "自然人失败成本、个人资产保护、债权追偿价值或相关主体暴露；需全文核对具体口径。"
    y = "待全文核对。"
    m = "待全文核对。"
    model = "待全文核对。"
    question = rec["title"]

    if "信贷" in category or "credit" in blob or "loan" in blob or "debt" in blob:
        y = "信贷供给、贷款条件、债务结构或融资可得性。"
        m = "债权追偿价值、贷款风险重新定价。"
    if "创业" in category or "entrepreneur" in blob or "startup" in blob or "innovation" in blob:
        y = "创业、风险承担、创新或企业进入。"
        m = "失败保险、重新开始、创业失败成本。"
    if "会计" in category or "audit" in blob or "earnings" in blob or "disclosure" in blob:
        y = "盈余管理、审计、披露或资本市场信息风险。"
        m = "财务忧虑、外部监督、坏消息隐藏或信息透明度。"
    if "股价崩盘" in rec["title"] or "crash" in blob:
        y = "股价崩盘风险/坏消息隐藏。"
        m = "坏消息隐藏、代理冲突、信息不透明。"

    return [
        rec["year"],
        apa_like(rec),
        category,
        x,
        exposure,
        y,
        m,
        model,
        question,
        rec["abstract"][:900],
        rec["relation"],
        rec["collision"],
        f"来自下载池初筛；原始文件：{rec['raw_file']}；建议：{rec['action']}",
    ]


def write_pool_workbook(records: list[dict]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    sheets = [
        ("高相关候选", [r for r in records if r["score"] >= 13]),
        ("疑似撞题", [r for r in records if "撞题" in r["collision"] or "强相邻" in r["collision"]]),
        ("中国知网候选", [r for r in records if r["source_db"] == "CNKI" and r["score"] >= 8]),
        ("Scopus高相关", [r for r in records if r["source_db"] == "Scopus" and r["score"] >= 8]),
        ("全部去重池", records),
    ]
    for name, subset in sheets:
        ws = wb.create_sheet(name)
        write_rows(ws, POOL_HEADERS, [pool_row(r) for r in subset])
        ws.auto_filter.ref = ws.dimensions
    wb.save(POOL_XLSX)


def write_review_v2(records: list[dict]) -> int:
    selected = []
    seen_titles = set()
    for rec in records:
        if rec["score"] < 13:
            continue
        if rec["title"] in seen_titles:
            continue
        seen_titles.add(rec["title"])
        selected.append(rec)
        if len(selected) >= 45:
            break

    wb = load_workbook(REVIEW_V1)
    ws = wb["文献综述"]
    existing_titles = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row and row[1]:
            existing_titles.add(str(row[1]).lower())
    appended = 0
    for rec in selected:
        title_key = rec["title"].lower()
        if any(title_key and title_key in old for old in existing_titles):
            continue
        ws.append(review_row_from_pool(rec))
        appended += 1
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    for idx, width in {
        1: 10,
        2: 48,
        3: 20,
        4: 30,
        5: 30,
        6: 26,
        7: 28,
        8: 24,
        9: 34,
        10: 62,
        11: 34,
        12: 30,
        13: 34,
    }.items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    wb.save(REVIEW_V2)
    return appended


def write_report(records: list[dict], appended: int) -> None:
    total = len(records)
    source_counts = Counter(r["source_db"] for r in records)
    category_counts = Counter(r["category"] for r in records)
    high = [r for r in records if r["score"] >= 13]
    suspicious = [r for r in records if "撞题" in r["collision"] or "强相邻" in r["collision"]]
    top = high[:20]

    lines = [
        "# 个人破产制度下载文献初筛报告",
        "",
        "更新日期：2026-05-04",
        "",
        "## 一、处理结果",
        "",
        f"- 去重后总记录：{total}",
        f"- Scopus：{source_counts.get('Scopus', 0)}",
        f"- CNKI：{source_counts.get('CNKI', 0)}",
        f"- 高相关候选（相关度分 >= 13）：{len(high)}",
        f"- 疑似撞题/强相邻：{len(suspicious)}",
        f"- 已追加到综述 v2 的候选行：{appended}",
        "",
        "## 二、类别分布",
        "",
    ]
    for category, count in category_counts.most_common():
        lines.append(f"- {category}：{count}")
    lines += [
        "",
        "## 三、最高优先级候选",
        "",
    ]
    for idx, rec in enumerate(top, 1):
        lines.append(f"{idx}. {rec['title']}（{rec['year']}，{rec['source_db']}，{rec['category']}，score={rec['score']}）")
        lines.append(f"   - 撞题判断：{rec['collision']}")
        lines.append(f"   - 建议动作：{rec['action']}")
    lines += [
        "",
        "## 四、当前判断",
        "",
        "这批材料确认了两个事实。第一，中国文献里已经有强相邻甚至直接覆盖的文章，尤其是《个人破产制度与企业盈余平滑》。第二，英文文献中个人破产、fresh start、创业失败成本、信贷供给和公司政策之间有一条成熟理论线，可以支撑我们把 X 从城市 PB dummy 改造成自然人责任边界冲击。",
        "",
        "后续不建议继续裸跑 `PB -> 某个企业Y`。更稳的写法是 `PB x 自然人责任暴露`，并把贡献写成同一制度冲击下不同自然人暴露企业的差异反应。",
        "",
        "## 五、产出文件",
        "",
        f"- 下载池初筛表：{POOL_XLSX}",
        f"- 文献综述 v2：{REVIEW_V2}",
    ]
    REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    records: list[dict] = []
    for path in RAW_DIR.glob("*.csv"):
        records.extend(parse_scopus_csv(path))
    for path in RAW_DIR.glob("*.txt"):
        records.extend(parse_cnki_txt(path))
    records = enrich(dedupe(records))
    write_pool_workbook(records)
    appended = write_review_v2(records)
    write_report(records, appended)
    print(f"records={len(records)}")
    print(f"pool={POOL_XLSX}")
    print(f"review_v2={REVIEW_V2}")
    print(f"report={REPORT}")
    print(f"appended={appended}")


if __name__ == "__main__":
    main()
