"""
机制分析章节重构：
1. 标题 "五、作用渠道分析" → "五、机制分析"
2. 三个渠道子项扩充理论论证+文献支撑（仿朱康2025风格）
3. 删除空段落 [105]-[115]
处理顺序：先替换文本，再从后往前删空段落
"""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy
import os

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC_PATH = f"{BASE}/manuscript/word/数据要素利用与资产定价效率v8.docx"

doc = Document(DOC_PATH)


def replace_para_text(para, new_text):
    """Clear all runs, add single new run with first run's formatting."""
    runs = para.runs
    first_rPr = None
    if runs:
        rPr_elem = runs[0]._element.find(qn('w:rPr'))
        if rPr_elem is not None:
            first_rPr = copy.deepcopy(rPr_elem)

    for r in runs:
        r._element.getparent().remove(r._element)
    for r_elem in list(para._element.findall(qn('w:r'))):
        para._element.remove(r_elem)

    new_run = OxmlElement('w:r')
    if first_rPr is not None:
        new_run.append(first_rPr)
    t_elem = OxmlElement('w:t')
    t_elem.text = new_text
    t_elem.set(qn('xml:space'), 'preserve')
    new_run.append(t_elem)
    para._element.append(new_run)


# ============================================================
# Step 1: 改标题
# ============================================================

p95 = doc.paragraphs[95]
assert '渠道' in p95.text or '机制' in p95.text, f"Expected 渠道/机制 at [95], got: {p95.text}"
replace_para_text(p95, "五、机制分析")
print(f"[95] heading → 五、机制分析")


# ============================================================
# Step 2: 替换三个渠道的正文段落
# ============================================================

# （一）信息中介渠道 [100]
analyst_text = (
    "分析师作为资本市场最重要的信息中介，通过深度调研、盈利预测和投资评级"
    "将企业层面的复杂信息转化为可供投资者使用的交易信号。"
    "Hong等（2000）发现分析师跟踪不足的股票信息融入价格的速度显著更慢，"
    "产生了可被动量策略利用的异常收益。"
    "黄俊和郭照蕊（2014）利用券商关闭这一准自然实验证实，"
    "外生的分析师覆盖减少导致股价延迟显著增加，"
    "为分析师覆盖与定价效率之间的因果关系提供了直接证据。"
    "数据要素利用涉及数据资产化、数字化转型和数据驱动决策等技术含量较高的业务实践，"
    "需要专业中介进行解读和传播。"
    "年报中数据利用信息的增加为分析师提供了新的研究素材和增量信息来源，"
    "降低了信息获取成本，从而可能吸引更多分析师关注和跟踪。"
    "带入机制模型的回归结果如表5第(1)(3)列所示，"
    "DU_kw对分析师覆盖（取对数）的同期系数为+0.0255（p<0.01），"
    "滞后一期系数为+0.0284（p<0.001），方向一致且系数在滞后口径下略有增大。"
    "经济意义上，DU_kw增加一个标准差（1.86）对应分析师跟踪人数提升约4.9%。"
    "进一步地，以t-1期分析师覆盖中位数分组，"
    "高覆盖组的DU_kw系数（-0.006，p<0.01）约为低覆盖组（-0.003，p<0.01）的1.9倍"
    "（Fisher p=0.074），"
    "这一梯度模式进一步验证了信息中介在数据信息传导中的作用。"
    "综上说明，数据要素利用可以通过提升分析师覆盖、"
    "拓展信息中介网络来降低股价延迟、提升定价效率。"
)

# （二）信息质量渠道 [102]
disp_text = (
    "分析师预测分歧度衡量的是市场参与者对企业价值判断的不一致程度，"
    "是信息环境质量的重要反映。"
    "较高的预测分歧意味着分析师所掌握的信息集存在显著差异，"
    "市场难以形成一致预期，价格发现过程受到阻碍。"
    "Zhang（2006）发现信息不确定性较高的企业预测分歧更大，"
    "且高分歧股票存在更严重的定价偏差。"
    "Diether等（2002）进一步证实，"
    "预测分歧度较高的股票未来收益率显著偏低。"
    "数据要素利用产生了更多结构化、可验证的运营信息，"
    "有助于降低分析师在预测企业未来盈利时面临的不确定性。"
    "与一般性文字叙述不同，数据驱动业务实践的披露往往伴随更具体的量化信息，"
    "为分析师提供了更硬的信息基础，"
    "有助于缩小分析师之间的信息差异，降低预测分歧度。"
    "带入机制模型的回归结果如表5第(2)(4)列所示，"
    "DU_kw对预测分歧度的同期系数为-0.0064（p<0.001），"
    "滞后一期系数为-0.0059（p<0.001），方向一致。"
    "DU_kw增加一个标准差对应分歧度降低约7.7%。"
    "Disp的显著下降同时也有助于区分两种竞争效应："
    "若数据利用主要带来信息模糊，分析师之间的理解差异应扩大而非缩小，"
    "实证结果与此方向相反。"
    "综上说明，数据要素利用可以通过降低分析师预测分歧、"
    "改善信息质量来促进价格发现。"
)

# （三）市场流动性渠道 [104]
amihud_text = (
    "信息要转化为定价效率的改善，最终需要通过市场交易来实现。"
    "Kyle（1985）和Glosten和Milgrom（1985）的经典模型表明，"
    "信息不对称程度较高时，做市商面临严重的逆向选择风险，"
    "倾向于扩大买卖价差以补偿潜在损失，"
    "由此产生的流动性不足阻碍了信息向价格的传递。"
    "Amihud和Mendelson（1986）在理论上证明了流动性溢价的存在，"
    "流动性较差的资产要求更高的预期回报作为补偿。"
    "Chordia等（2008）从实证上发现，流动性改善加速了信息融入价格的过程，"
    "为流动性与定价效率之间的正向关联提供了经验证据。"
    "数据要素利用通过增加企业可观测信息的供给，"
    "缩小内外部信息不对称程度，可能缓解逆向选择问题，"
    "降低交易摩擦从而改善市场流动性。"
    "与前两条渠道侧重信息的生产和加工不同，"
    "市场流动性渠道刻画的是信息融入价格的交易环节，"
    "三者共同构成了从信息供给到定价效率改善的完整链条。"
    "带入机制模型的回归结果如表5第(5)(6)列所示，"
    "DU_kw对Amihud非流动性指标的同期系数为-0.0010（p<0.01），"
    "滞后一期系数为-0.0014（p<0.001），方向一致且滞后口径系数绝对值更大。"
    "DU_kw增加一个标准差对应Amihud降低约3.9%。"
    "综上说明，数据要素利用可以通过改善市场流动性、"
    "降低交易摩擦来加速信息融入价格。"
)

# Apply replacements
for idx, new_text, label in [
    (100, analyst_text, "信息中介渠道"),
    (102, disp_text, "信息质量渠道"),
    (104, amihud_text, "市场流动性渠道"),
]:
    p = doc.paragraphs[idx]
    old_preview = p.text[:50]
    replace_para_text(p, new_text)
    print(f"[{idx}] {label}: {old_preview}... → {new_text[:50]}...")


# ============================================================
# Step 3: 删除空段落 [105]-[115]（从后往前）
# ============================================================

deleted = 0
for i in range(115, 104, -1):
    p = doc.paragraphs[i]
    if not p.text.strip():
        p._element.getparent().remove(p._element)
        deleted += 1
        print(f"  Deleted empty paragraph [{i}]")
    else:
        print(f"  SKIPPED [{i}] (not empty): {p.text[:40]}")

print(f"\nDeleted {deleted} empty paragraphs")


# ============================================================
# Step 4: 验证
# ============================================================
print("\n=== 验证最终结构 ===")
for i, p in enumerate(doc.paragraphs):
    if p.style.name.startswith('Heading') and '机制' in p.text:
        start = i
        break

for j in range(start, min(start + 16, len(doc.paragraphs))):
    pp = doc.paragraphs[j]
    style = pp.style.name
    text = pp.text[:90] if pp.text else "(empty)"
    print(f"[{j}] {style}: {text}")

doc.save(DOC_PATH)
print(f"\nDone! Saved to {DOC_PATH}")
