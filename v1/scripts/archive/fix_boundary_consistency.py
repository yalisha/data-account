"""
异质性边界条件口径统一：
[1] Para [81]: 设计段移除分析师覆盖，加入市场流动性
[2] Para [43]: H3 → 行业属性理论动机
[3] Para [42]: 末尾"据此提出："→ 渠道梯度检验引导
[4] Para [40]: 两个维度 → 三个维度，分析师覆盖定位为渠道预测
从高索引往低索引改，避免索引偏移
"""
from docx import Document
import os

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC_PATH = f"{BASE}/manuscript/word/数据要素利用与资产定价效率v7.docx"
doc = Document(DOC_PATH)


def replace_para(idx, new_text):
    p = doc.paragraphs[idx]
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    print(f"  Para [{idx}] updated ({len(new_text)} chars)")


# ============================================================
# [1] Para [81]: 设计段 — 分析师覆盖→市场流动性
# ============================================================
new_81 = (
    "异质性分析方面，分别按产权性质、"
    "市场流动性（Amihud中位数分割）、"
    "企业规模（年度中位数分割）和行业属性（高科技vs传统）进行分组回归，"
    "采用Fisher组间差异检验评估系数差异的统计显著性。"
)
print("[1] Updating Para [81] (design: analyst→liquidity)...")
replace_para(81, new_81)


# ============================================================
# [2] Para [43]: H3 → 行业属性理论动机
# ============================================================
new_43 = (
    "行业属性也可能调节效应的强度。"
    "传统行业的信息环境基线透明度较低，"
    "市场参与者对企业数据利用行为的认知缺口更大，"
    "增量信息的边际价值因此更高，"
    "定价效率的改善空间相应更为充裕。"
)
print("[2] Updating Para [43] (H3 → industry theory)...")
replace_para(43, new_43)


# ============================================================
# [3] Para [42]: 末尾"据此提出："→ 渠道梯度检验引导
# ============================================================
old_42 = doc.paragraphs[42].text
new_42 = old_42.replace(
    "据此提出：",
    "这一梯度含义在渠道分析中通过t-1期分析师覆盖的分组检验加以考察。"
)
if new_42 == old_42:
    print("  WARNING: '据此提出：' not found in [42]!")
else:
    print("[3] Updating Para [42] (remove '据此提出')...")
    replace_para(42, new_42)


# ============================================================
# [4] Para [40]: 两个维度 → 三个维度 + 分析师覆盖重新定位
# ============================================================
old_40 = doc.paragraphs[40].text
# 替换核心句
new_40 = old_40.replace(
    "从两个维度检验效应的边界条件："
    "产权性质反映企业信息环境的制度基础，"
    "市场流动性反映信息融入价格的物理约束。"
    "此外，分析师覆盖维度基于渠道检验的理论预期形成正式假说。",

    "从三个维度检验效应的边界条件："
    "产权性质反映企业信息环境的制度基础，"
    "市场流动性反映信息融入价格的物理约束，"
    "行业属性反映信息环境的基线透明度差异。"
    "此外，渠道检验预测了分析师覆盖的梯度含义，在机制分析中加以考察。"
)
if new_40 == old_40:
    print("  WARNING: replacement text not found in [40]!")
else:
    print("[4] Updating Para [40] (2 dims → 3 dims)...")
    replace_para(40, new_40)


# Save
doc.save(DOC_PATH)
print(f"\nDone! 4 paragraphs updated in v7.docx")
