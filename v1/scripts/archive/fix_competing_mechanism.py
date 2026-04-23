"""
竞争机制区分：从"看系数符号"到"差异化预测"
用两条已有证据（Disp下降 + 投资组合非对称）区分供给效应与模糊效应
从高索引往低索引改，避免插入操作的索引偏移
"""
from docx import Document
from copy import deepcopy
import os

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC_PATH = f"{BASE}/manuscript/word/数据要素利用与资产定价效率v7.docx"
doc = Document(DOC_PATH)


def replace_para(idx, new_text):
    p = doc.paragraphs[idx]
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    print(f"  Para [{idx}] replaced ({len(new_text)} chars)")


def insert_para_after(idx, new_text):
    """在 idx 段落之后插入新段落，继承其格式"""
    ref_para = doc.paragraphs[idx]
    new_para = deepcopy(ref_para._element)
    # 清空内容
    for child in list(new_para):
        if child.tag.endswith('}r'):
            new_para.remove(child)
    ref_para._element.addnext(new_para)
    # 重新获取新段落并写入文本
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = new_text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    new_para.append(r)
    print(f"  Inserted after [{idx}] ({len(new_text)} chars)")


# ============================================================
# Step 1: Para [126] 结论 — 在"信息供给效应占据主导地位"后加差异化证据概括
# ============================================================
new_126 = (
    "本文以2011至2024年沪深A股上市公司为样本，"
    "用年报全文五维度关键词体系度量企业的数据要素利用强度，"
    "考察其对资本市场定价效率的影响。"
    "该指标测度的是年报文本所反映的数据利用行为，"
    "区别于数据资产披露决策和更宽泛的数字化转型状态。"
    "研究发现，企业数据要素利用强度的提升降低了股价延迟，"
    "即提升了资本定价效率，信息供给效应占据主导地位。"
    "这一判断不仅基于基准回归中的净效应方向，"
    "还得到差异化含义的支持："
    "分析师预测分歧度随数据利用增加而显著降低，"
    "与信息模糊效应中分歧应扩大的预期相反；"
    "投资组合中高数据利用企业回归合理估值而非遭受定价惩罚，"
    "排除了模糊效应对高利用企业的实质约束。"
    "滞后一期回归、披露后窗口回归和前导项检验均支持主结论的时序稳健性，"
    "工具变量估计与之方向一致，提供了补充性支持。"
    "渠道分析采用江艇（2022）两步法，"
    "提供了与数据利用经由外部信息中介影响定价效率一致的证据。"
    "信息中介渠道方面，数据利用与分析师跟踪正相关，滞后一期方向一致；"
    "信息质量渠道方面，数据利用降低了预测分歧度，时序证据一致。"
    "按t-1期分析师覆盖分组的梯度检验进一步提供了与渠道解释一致的模式。"
    "这一外部信息传导机制与既有文献所关注的企业内部行为渠道形成互补。"
    "效应在不同产权性质和流动性环境下普遍存在，"
    "市场流动性和行业属性是两个显著的边界条件。"
    "投资组合层面，数据利用的定价含义主要表现为"
    "低数据利用企业承受信息不透明折价，"
    "而非高数据利用企业获得系统性溢价，"
    "这一非对称模式与信息摩擦减少假说一致。"
)
print("Step 1: Updating Para [126] (conclusion)...")
replace_para(126, new_126)


# ============================================================
# Step 2: 在 Para [117] 后插入区分论证段落
# ============================================================
distinction_para = (
    "上述渠道检验和投资组合检验的结果，"
    "除了为信息供给效应提供传导路径的证据外，"
    "还从差异化预测的角度将信息供给效应与信息模糊效应加以区分。"
    "第一，分析师预测分歧度的显著下降"
    "（Disp系数-0.006，p<0.001）"
    "与信息模糊效应的核心预测方向相反。"
    "若数据利用带来的信息以模糊性为主要特征，"
    "分析师之间的理解差异应当扩大而非缩小。"
    "分歧度的降低表明，数据利用传递的增量信息"
    "在缩小而非扩大分析师群体的认知差异。"
    "第二，投资组合检验中的非对称结构提供了补充证据。"
    "Q1（最低利用组）在FF3下承受了显著的负异常收益"
    "（-0.28%，t=-2.08），"
    "而Q5（最高利用组）的异常收益在统计上为零"
    "（+0.09%，t=0.38）。"
    "若信息模糊效应对高数据利用企业构成实质性的定价惩罚，"
    "Q5组合也应出现负向偏离，"
    "但高利用企业恰恰回归了合理估值。"
    "两项证据分别从信息质量和截面收益的维度"
    "指向信息供给效应占优的一致结论。"
)
print("Step 2: Inserting distinction paragraph after [117]...")
insert_para_after(117, distinction_para)


# ============================================================
# Step 3: Para [88] 基准回归 — 在"H1得到支持"后加前向引导
# ============================================================
old_88 = doc.paragraphs[88].text
# 在末尾追加
append_88 = (
    "基准回归确认了净效应的方向，"
    "但仅凭系数符号尚不能区分信息供给效应确实占优还是两种效应并存而供给碰巧更强。"
    "后文的渠道检验和投资组合分析将从差异化含义的角度提供进一步的区分证据。"
)
new_88 = old_88 + append_88
print("Step 3: Appending to Para [88] (baseline results)...")
replace_para(88, new_88)


# ============================================================
# Step 4: Para [29] 理论部分 — 在"据此提出假说H1"前加差异化预测
# ============================================================
old_29 = doc.paragraphs[29].text
# 在"据此提出假说H1"前插入
insert_text = (
    "除了净效应方向外，两种效应还生成差异化的可观测含义。"
    "若信息供给效应占优，分析师预测分歧度应随数据利用的增加而降低，"
    "因为更充分的可信信息缩小了分析师之间的信息差异；"
    "反之，若信息模糊效应占优，分歧度应上升。"
    "类似地，若信息模糊效应对高数据利用企业构成实质性约束，"
    "高利用企业在投资组合中也应表现出定价惩罚而非回归合理估值。"
    "后续实证将在确认净方向之后，通过这些差异化含义提供进一步的区分证据。"
)
new_29 = old_29.replace("据此提出假说H1", insert_text + "据此提出假说H1")
if new_29 == old_29:
    print("  WARNING: '据此提出假说H1' not found in [29]!")
else:
    print("Step 4: Inserting differential predictions into Para [29]...")
    replace_para(29, new_29)


# ============================================================
# Step 5: Para [9] 摘要 — 在"信息供给效应占主导"后加一句
# ============================================================
old_9 = doc.paragraphs[9].text
insert_after = "信息供给效应占主导。"
add_9 = (
    "分析师预测分歧度的显著下降和高数据利用企业回归合理估值的非对称模式，"
    "从差异化含义角度提供了区分两种效应的证据。"
)
new_9 = old_9.replace(insert_after, insert_after + add_9)
if new_9 == old_9:
    print("  WARNING: insertion point not found in [9]!")
else:
    print("Step 5: Updating Para [9] (abstract)...")
    replace_para(9, new_9)


# ============================================================
# Step 6: Para [2] Highlights — 在"信息供给效应占优"后加半句
# ============================================================
old_2 = doc.paragraphs[2].text
insert_after_2 = "信息供给效应占优。"
add_2 = (
    "分析师预测分歧度的降低和投资组合非对称定价模式"
    "从差异化含义角度进一步支持了这一判断。"
)
new_2 = old_2.replace(insert_after_2, insert_after_2 + add_2)
if new_2 == old_2:
    print("  WARNING: insertion point not found in [2]!")
else:
    print("Step 6: Updating Para [2] (Highlights)...")
    replace_para(2, new_2)


# Save
doc.save(DOC_PATH)
print(f"\nDone! 5 paragraphs modified + 1 paragraph inserted in v7.docx")
