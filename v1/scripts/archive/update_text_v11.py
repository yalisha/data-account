"""
更新v7.docx正文：扩展渠道分析和异质性分析的文字
"""

from docx import Document
from docx.shared import Pt
from copy import deepcopy

BASE = "/Users/mac/computerscience/15会计研究"
DOCX = f"{BASE}/manuscript/word/数据要素利用与资产定价效率v7.docx"

doc = Document(DOCX)


def replace_para_text(idx, old_text_start, new_text):
    """Replace paragraph text, preserving first run's formatting."""
    p = doc.paragraphs[idx]
    current = p.text
    if not current.startswith(old_text_start[:30]):
        print(f"  WARNING: Para [{idx}] does not start with expected text!")
        print(f"  Expected: {old_text_start[:60]}...")
        print(f"  Got:      {current[:60]}...")
        return False

    # Preserve formatting from first run
    if p.runs:
        fmt_run = p.runs[0]
        font_name = fmt_run.font.name
        font_size = fmt_run.font.size
        bold = fmt_run.bold
    else:
        font_name = None
        font_size = None
        bold = None

    # Clear all runs
    for run in p.runs:
        run.text = ''
    # Set first run or add new
    if p.runs:
        p.runs[0].text = new_text
    else:
        run = p.add_run(new_text)
        if font_name:
            run.font.name = font_name
        if font_size:
            run.font.size = font_size
        if bold is not None:
            run.bold = bold
    return True


def insert_para_after(idx, text):
    """Insert a new paragraph after paragraph idx, copying format from idx."""
    ref_para = doc.paragraphs[idx]
    # Create new paragraph element by copying reference
    new_p = deepcopy(ref_para._element)
    # Clear text in copy
    for r in new_p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r'):
        for t in r.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
            t.text = ''
    # Set text
    runs = new_p.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r')
    if runs:
        ts = runs[0].findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')
        if ts:
            ts[0].text = text
        else:
            from lxml import etree
            ns = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
            t_elem = etree.SubElement(runs[0], f'{{{ns}}}t')
            t_elem.text = text
    ref_para._element.addnext(new_p)
    return new_p


# ============================================================
# 1. Update Section 五 (渠道分析)
# ============================================================
print("Updating Section 五...")

# [106] Update intro to mention 3 dimensions
replace_para_text(106,
    "总效应已经确认，传导路径是什么？以下从信息中介和信息质量两个维度",
    "总效应已经确认，传导路径是什么？以下从信息中介、信息质量和资产配置行为三个维度，用同期与滞后一期回归加以检验。"
)

# [108] Update Analyst paragraph - fix lagged coef to match new results
replace_para_text(108,
    "年报中关于数据驱动战略和数字化转型的内容",
    "年报中关于数据驱动战略和数字化转型的内容，因其信息技术含量，更容易吸引专业分析师的注意。同期检验结果显示，DU_kw对分析师覆盖（取对数）的系数为+0.0255（t=3.15，p=0.002）。滞后一期回归中，DU_kw(t)对Analyst(t+1)的系数为+0.0284（t=3.70，p<0.001），系数绝对值反而略有增大，支持了数据要素利用对分析师覆盖的时序因果关系。经济意义方面，DU_kw增加一个标准差（1.86）对应分析师覆盖对数增加0.047，相当于分析师跟踪人数提升约4.9%。第二步，分析师覆盖对股价延迟的降低作用已在大量文献中得到因果层面的确认：Hong等（2000）发现分析师跟踪不足的股票存在动量利润，黄俊和郭照蕊（2014）利用券商关闭事件证实分析师覆盖减少会增加股价延迟。两步证据合在一起，分析师扮演了数据信息传导至股价的中介角色。"
)

# [109] Update Disp paragraph - fix lagged coef
replace_para_text(109,
    "第二条渠道关注预测分歧度",
    "第二条渠道关注预测分歧度。DU_kw对Disp的同期回归系数为-0.0064（t=-4.25，p<0.001）。滞后一期回归中，DU_kw(t)对Disp(t+1)的系数为-0.0059（t=-3.44，p<0.001），方向一致，排除了同期共同决定的干扰。经济意义方面，DU_kw增加一个标准差对应预测分歧度降低0.012，约为样本均值的7.7%。预测分歧度是信息不对称的代理变量：Zhang（2006）证实预测分歧度较高的股票具有更大的定价偏差和更低的价格效率，分歧度的降低有助于市场对企业价值形成更一致的判断，加速价格发现过程。Analyst和Disp分别对应信息生产的广度和深度，两条路径互补。"
)

# Insert new paragraph for FinAsset channel after [109]
print("  Inserting FinAsset channel paragraph...")
finasset_text = (
    "第三条渠道考察金融资产配置行为。数据要素利用水平较高的企业在资本市场中往往具备更强的信息获取和处理能力，从而更倾向于配置金融资产以提升资金运营效率（朱康和唐勇，2025；李姝等，2025）。"
    "同期回归中，DU_kw对FinAsset的系数为+0.0008（t=2.46，p=0.014），滞后一期回归系数为+0.0010（t=2.58，p=0.010），同期与滞后结果方向一致、均在5%水平上通过检验。"
    "金融资产配置比例的提高意味着企业与资本市场之间建立了更频繁的交易联系，市场参与者可以通过观察企业的金融市场行为获取增量信息，有助于降低信息不对称、提升定价效率（Shao等，2024）。"
)
insert_para_after(109, finasset_text)
# After insertion, paragraph indices shift by 1 for all subsequent paragraphs

# [110→111] Update summary - now references 3 channels
replace_para_text(111,
    "两条渠道在同期和滞后一期回归中均保持高度显著",
    "三条渠道在同期和滞后一期回归中均保持显著且系数方向一致，满足江艇（2022）关于渠道检验时序稳健性的要求。Analyst的N为43,843和37,934，Disp的N为25,092和21,424，FinAsset的N为43,843和37,892。Disp的样本量较小是因为分歧度指标仅在有至少两位分析师预测时才可计算。"
)

# [111→112] Update table note
replace_para_text(112,
    "注：遵循江艇（2022）两步法",
    "注：遵循江艇（2022）两步法，Panel A报告同期回归、Panel B报告滞后一期回归。固定效应为企业+年份，标准误按行业×年份聚类。Analyst为分析师覆盖对数，Disp为分析师预测分歧度，FinAsset为金融资产占比。M→Y的理论关系参见正文讨论。"
)

# [112→113] Update conclusion paragraph
replace_para_text(113,
    "渠道检验的结论可以概括为",
    "渠道检验的结论可以概括为：信息中介渠道方面，数据利用通过吸引更多分析师跟踪而加速信息融入价格，同期与滞后一期回归方向一致、统计可靠。信息质量渠道方面，数据利用降低了预测分歧度，滞后一期结果同样稳健，支持时序因果解释。资产配置行为渠道方面，数据利用提升了金融资产配置比例，增加了企业与资本市场的信息交互频率。三条渠道分别从信息供给广度、信息精度和市场参与行为三个层面为总效应提供了传导解释。"
)

# ============================================================
# 2. Update Section 六 (异质性分析)
# ============================================================
print("Updating Section 六...")

# After insertion above, indices shift by 1
# Original [113]→[114] 六、异质性分析
# Original [114]→[115] intro
# Original [115]→[116] Panel A
# Original [116]→[117] Panel B
# Original [117]→[118] conclusion
# Original [118]→[119] note

# [115] Update intro
replace_para_text(115,
    "传导机制确认后，下一步看效应是否因企业特征而异",
    "传导机制确认后，下一步看效应是否因企业特征而异。交互项回归检验两个边界条件（产权性质和市场流动性），分组回归从产权性质、信息中介活跃度、企业规模和行业属性四个维度描述效应的异质性梯度。"
)

# [117] Expand Panel B text to include Size and HighTech
replace_para_text(117,
    "Panel B的分组回归提供了描述性补充",
    "Panel B的分组回归提供了描述性补充。产权性质方面，国有企业样本中DU_kw的系数为-0.006（p<0.01，N=14,412），民营企业样本中为-0.003（p<0.01，N=29,369），两组均在1%水平上通过检验，国企效应约为民企的1.8倍。Fisher z检验的p值为0.094，在10%水平上边际拒绝系数相等的原假设。信息中介活跃度方面，高分析师覆盖企业的系数为-0.006（p<0.01，N=21,836），低覆盖企业为-0.003（p<0.01，N=21,457），高覆盖组效应约为低覆盖组的1.7倍，与渠道检验中分析师覆盖的结果互为印证。"
)

# Insert new paragraph for Size and HighTech after [117]
size_tech_text = (
    "企业规模方面，大企业的系数为-0.005（p<0.01，N=21,463），小企业为-0.005（p<0.01，N=21,578），两组系数几乎相同（Fisher p=0.966），说明数据要素利用的定价效率改善作用不受企业规模制约。"
    "行业属性方面，高科技行业（信息传输+计算机通信电子）的系数为-0.002（p<0.05，N=7,473），传统行业的系数为-0.004（p<0.01，N=36,328），Fisher z检验的p值为0.055，在10%水平上边际显著。传统行业效应更强的结果与理论预期一致：传统行业的信息环境相对模糊，数据要素利用带来的增量信息价值更大，定价效率的改善空间也更大。"
)
insert_para_after(117, size_tech_text)

# [118→119] Update conclusion
replace_para_text(119,
    "异质性分析的结论比较清楚",
    "异质性分析的结论比较清楚。数据要素利用对定价效率的改善在所有子样本中均成立，不存在效应反转的情形。两个通过统计检验的边界条件是市场流动性和行业属性。流动性差的股票，数据信息融入价格的摩擦更大，效应因此减弱；传统行业由于信息环境相对模糊，数据利用的边际改善更为明显。此外，产权性质的调节效应在控制流动性后边际成立，而企业规模对效应无差异化影响。"
)

# [119→120] Update note
replace_para_text(120,
    "注：Panel A报告全样本交互项回归",
    "注：Panel A报告全样本交互项回归，Amihud_c为中心化后的Amihud非流动性指标（减去样本均值）。Panel B报告分组回归结果，高/低分析师覆盖和大/小企业的分割点为全样本中位数，高科技行业包括信息传输业和计算机通信电子制造业（证监会行业分类I类和C39类）。所有模型均采用企业+年份双向固定效应，标准误在行业×年份水平上聚类。"
)


# ============================================================
# 3. Update table number references in text
# ============================================================
print("Updating table references...")

# Para 101 (portfolio): "表3的结果" → should be 附表1
# After +2 insertions above, original 101 is now 101 (insertions were after 109 and 117)
# Actually wait - insertions at 109 and 117 shift paragraphs AFTER those indices
# Para 101 is BEFORE the insertions, so it stays at 101
p101 = doc.paragraphs[101]
old_text = p101.text
if '表3的结果' in old_text:
    new_text = old_text.replace('表3的结果', '附表1的结果')
    # Simple text replacement preserving formatting
    for run in p101.runs:
        if '表3' in run.text:
            run.text = run.text.replace('表3的结果', '附表1的结果')
    print("  Updated '表3的结果' -> '附表1的结果' in para 101")


# ============================================================
# Save
# ============================================================
print("Saving...")
doc.save(DOCX)
print("Done!")

# Count new section lengths
doc2 = Document(DOCX)
channel_chars = sum(len(doc2.paragraphs[i].text) for i in range(105, 114))
het_chars = sum(len(doc2.paragraphs[i].text) for i in range(114, 121))
print(f"\nNew section lengths:")
print(f"  渠道分析: {channel_chars}字 (was ~1352)")
print(f"  异质性: {het_chars}字 (was ~945)")
