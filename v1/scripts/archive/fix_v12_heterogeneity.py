from docx import Document

doc_path = "/Users/mac/computerscience/15会计研究/manuscript/word/数据要素利用与资产定价效率v8.docx"
doc = Document(doc_path)

def replace_in_para(idx, old, new):
    p = doc.paragraphs[idx]
    if old not in p.text:
        print(f"  [WARNING] [{idx}] NOT FOUND: '{old}'")
        return False
    for run in p.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            print(f"  [{idx}] replaced: '{old}' -> '{new}'")
            return True
    # concat fallback
    combined = ''.join(r.text for r in p.runs)
    if old in combined:
        new_text = combined.replace(old, new)
        for r in p.runs[1:]:
            p._element.remove(r._element)
        p.runs[0].text = new_text
        print(f"  [{idx}] concat-replaced: '{old}' -> '{new}'")
        return True
    print(f"  [WARNING] [{idx}] run fallback failed")
    return False

print("=== 更新异质性分析系数 ===")

# [108] SOE section
# 非国企: -0.003 -> -0.004
replace_in_para(108, "非国有企业组（-0.003）的1.8倍", "非国有企业组（-0.004）的1.6倍")
# Fisher p: 0.094 -> 0.106
replace_in_para(108, "P值为0.094", "P值为0.106")
# 0.094意味着10%显著，0.106意味着不显著
replace_in_para(108, "说明组间差异性显著存在。原因可能在于", "虽未通过10%显著性检验，但差异方向与理论预期一致。原因可能在于")

# [110] Liquidity section
# 低流动性: -0.003 -> -0.004
replace_in_para(110, "低流动性组为-0.003", "低流动性组为-0.004")
# Fisher p: 0.805 -> 0.868
replace_in_para(110, "P值为0.805", "P值为0.868")
# 交互项: 0.025 -> 0.026
replace_in_para(110, "交互项系数为0.025", "交互项系数为0.026")

# [114] Size section
# Fisher p: 0.966 -> 0.695
replace_in_para(114, "P值为0.966", "P值为0.695")

# [116] Industry section
# 传统行业: -0.004 -> -0.005
replace_in_para(116, "传统行业为-0.004（p<0.01，N=36,328）", "传统行业为-0.005（p<0.01，N=36,328）")
# Fisher p: 0.055 -> 0.025
replace_in_para(116, "P值为0.055", "P值为0.025")

doc.save(doc_path)
print("\n=== Done! ===")

# Verify key changes
doc2 = Document(doc_path)
for i in [108, 110, 114, 116]:
    t = doc2.paragraphs[i].text
    # Check for key values
    print(f"\n[{i}]: ...{t[t.find('费舍尔')-5:t.find('费舍尔')+40] if '费舍尔' in t else t[:100]}...")
