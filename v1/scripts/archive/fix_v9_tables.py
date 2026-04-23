"""
修复v9.docx中的三个表格问题：
1. Table 5（机制分析）：加Amihud列，更新全部数值匹配正文
2. Table 6（异质性）：(3)(4)列标题从"分析师覆盖"改为"流动性"，更新数值
3. Table 0（变量定义）：分离控制变量和机制变量，补Disp行

核心策略：Table 5 用XML直接操作tc元素，不依赖table.cell()索引
"""

from docx import Document
from copy import deepcopy
from docx.oxml.ns import qn

doc_path = "/Users/mac/computerscience/15会计研究/manuscript/word/数据要素利用与资产定价效率v9.docx"
doc = Document(doc_path)


def get_tcs(table, row_idx):
    """获取指定行的所有tc元素"""
    return table.rows[row_idx]._tr.findall(qn('w:tc'))


def set_tc_text(tc, text):
    """直接设置tc元素的文字（保留格式）"""
    paragraphs = tc.findall(qn('w:p'))
    if not paragraphs:
        return
    p = paragraphs[0]
    runs = p.findall(qn('w:r'))
    if runs:
        # 保留第一个run的格式，清除其他runs
        for r in runs[1:]:
            p.remove(r)
        # 设置第一个run的文字
        t_elem = runs[0].find(qn('w:t'))
        if t_elem is not None:
            t_elem.text = text
            # 保持空格
            t_elem.set(qn('xml:space'), 'preserve')
        else:
            t_elem = runs[0].makeelement(qn('w:t'), {})
            t_elem.text = text
            t_elem.set(qn('xml:space'), 'preserve')
            runs[0].append(t_elem)
    else:
        # 没有runs，创建一个run
        from docx.oxml import OxmlElement
        r_elem = OxmlElement('w:r')
        t_elem = OxmlElement('w:t')
        t_elem.text = text
        t_elem.set(qn('xml:space'), 'preserve')
        r_elem.append(t_elem)
        p.append(r_elem)


def set_cell_text(table, row_idx, col_idx, text):
    """通过table.cell()设置文字（仅用于未修改列数的表格）"""
    cell = table.cell(row_idx, col_idx)
    for p in cell.paragraphs:
        for run in p.runs:
            run.text = ""
    if cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = text
    else:
        cell.paragraphs[0].text = text


# ============================================================
# Fix 1: Table 5 - 机制分析表
# 从5列(1 label + 4 data)扩展到7列(1 label + 6 data)，加Amihud同期+滞后
# ============================================================
print("=== Fix 1: Table 5 (机制分析) ===")
t5 = doc.tables[5]
n_rows = len(t5.rows)
print(f"  行数: {n_rows}")

# Step 1: 给每行加2个tc元素，并更新tblGrid
# 先更新tblGrid
tbl = t5._tbl
tblGrid = tbl.find(qn('w:tblGrid'))
grid_cols = tblGrid.findall(qn('w:gridCol'))
print(f"  原始gridCol数: {len(grid_cols)}")
# 克隆最后一个gridCol两次
for _ in range(2):
    new_gc = deepcopy(grid_cols[-1])
    tblGrid.append(new_gc)
print(f"  新gridCol数: {len(tblGrid.findall(qn('w:gridCol')))}")

# Step 2: 给每行加2个新tc
for row in t5.rows:
    tr = row._tr
    tcs = tr.findall(qn('w:tc'))
    last_tc = tcs[-1]
    for _ in range(2):
        new_tc = deepcopy(last_tc)
        # 清空文字
        for p in new_tc.findall(qn('w:p')):
            for r in p.findall(qn('w:r')):
                t_el = r.find(qn('w:t'))
                if t_el is not None:
                    t_el.text = ''
        tr.append(new_tc)

# 验证列数
print(f"  扩展后列数: {len(t5.columns)}")

# Step 3: 写入数据（通过XML直接操作）
mech_data = {
    'Analyst': {"coef": 0.0273, "se": 0.0082, "N": 43843, "R2": 0.793},
    'Disp':    {"coef": -0.0059, "se": 0.0015, "N": 25092, "R2": 0.482},
    'Amihud':  {"coef": -0.0010, "se": 0.0003, "N": 43843, "R2": 0.562},
}
mech_lagged = {
    'Analyst': {"coef": 0.0335, "se": 0.0085, "N": 37404, "R2": 0.816},
    'Disp':    {"coef": -0.0064, "se": 0.0018, "N": 21160, "R2": 0.492},
    'Amihud':  {"coef": -0.0014, "se": 0.0003, "N": 37404, "R2": 0.629},
}

def fmt_coef(v, sig="***"):
    return f"{v:.4f}{sig}"
def fmt_se(v):
    return f"({v:.4f})"
def fmt_n(v):
    return f"{v:,}"
def fmt_r2(v):
    return f"{v:.3f}"

vars_order = ['Analyst', 'Disp', 'Amihud']

# 定义每行7列的内容
all_row_data = [
    # Row 0: 列号
    ['', '(1)', '(2)', '(3)', '(4)', '(5)', '(6)'],
    # Row 1: DV名
    ['被解释变量', 'Analystt', 'Dispt', 'Amihudt', 'Analystt+1', 'Dispt+1', 'Amihudt+1'],
    # Row 2: Panel A标题
    ['Panel A: 同期回归', '', '', '', '', '', ''],
    # Row 3: DUkw系数 Panel A (cols 1-3)
    ['DUkw'] + [fmt_coef(mech_data[v]['coef']) for v in vars_order] + ['', '', ''],
    # Row 4: SE Panel A
    [''] + [fmt_se(mech_data[v]['se']) for v in vars_order] + ['', '', ''],
    # Row 5: N Panel A
    ['N'] + [fmt_n(mech_data[v]['N']) for v in vars_order] + ['', '', ''],
    # Row 6: R² Panel A
    ['R²'] + [fmt_r2(mech_data[v]['R2']) for v in vars_order] + ['', '', ''],
    # Row 7: Panel B标题
    ['Panel B: 滞后一期', '', '', '', '', '', ''],
    # Row 8: DUkw系数 Panel B (cols 4-6)
    ['DUkw', '', '', ''] + [fmt_coef(mech_lagged[v]['coef']) for v in vars_order],
    # Row 9: SE Panel B
    ['', '', '', ''] + [fmt_se(mech_lagged[v]['se']) for v in vars_order],
    # Row 10: N Panel B
    ['N', '', '', ''] + [fmt_n(mech_lagged[v]['N']) for v in vars_order],
    # Row 11: R² Panel B
    ['R²', '', '', ''] + [fmt_r2(mech_lagged[v]['R2']) for v in vars_order],
    # Row 12: Controls
    ['Controls'] + ['YES'] * 6,
    # Row 13: FE
    ['Firm/Year FE'] + ['YES'] * 6,
]

for i, row_data in enumerate(all_row_data):
    tcs = get_tcs(t5, i)
    for j, val in enumerate(row_data):
        if j < len(tcs):
            set_tc_text(tcs[j], val)
    print(f"  Row {i}: {row_data}")

print("  Table 5 done!")

# ============================================================
# Fix 2: Table 6 - 异质性分析表
# (3)(4)列从"高/低分析师覆盖"改为"低/高流动性"，更新数值
# ============================================================
print("\n=== Fix 2: Table 6 (异质性) ===")
t6 = doc.tables[6]

# Panel B header row (row 14)
set_cell_text(t6, 14, 3, '(3) 低流动性')
set_cell_text(t6, 14, 4, '(4) 高流动性')
print("  Fixed column headers: (3) 低流动性, (4) 高流动性")

# Panel B DUkw coefficients (row 16) - update cols 3,4
# generate_tables_v11.py:
# 低流动性: coef=-0.00361, se=0.00100, N=21448, R²=0.507
# 高流动性: coef=-0.00386, se=0.00115, N=21384, R²=0.432
set_cell_text(t6, 16, 3, '-0.0036***')
set_cell_text(t6, 16, 4, '-0.0039***')
print("  Fixed coefficients: -0.0036***, -0.0039***")

# SE row (row 17)
set_cell_text(t6, 17, 3, '(0.0010)')
set_cell_text(t6, 17, 4, '(0.0012)')
print("  Fixed SEs")

# N row (row 20)
set_cell_text(t6, 20, 3, '21,448')
set_cell_text(t6, 20, 4, '21,384')
print("  Fixed Ns")

# R² row (row 21)
set_cell_text(t6, 21, 3, '0.507')
set_cell_text(t6, 21, 4, '0.432')
print("  Fixed R²s")

print("  Table 6 done!")

# ============================================================
# Fix 3: Table 0 - 变量定义表
# 分离控制变量和机制变量
# ============================================================
print("\n=== Fix 3: Table 0 (变量定义) ===")
t0 = doc.tables[0]

# Step 1: 交换 row 13 (Amihud) 和 row 15 (AuditType) 的文字内容
r13_data = [t0.cell(13, j).text for j in range(4)]
r15_data = [t0.cell(15, j).text for j in range(4)]
print(f"  Swapping row 13 ({r13_data[0]}) <-> row 15 ({r15_data[0]})")

for j in range(4):
    set_cell_text(t0, 13, j, r15_data[j])
    set_cell_text(t0, 15, j, r13_data[j])
# 交换后：Row 12=InstHold, Row 13=AuditType, Row 14=Analyst, Row 15=Amihud

# Step 2: 在Analyst行前插入分隔行 "机制变量"
last_row = t0.rows[-1]
sep_tr = deepcopy(last_row._tr)
tcs = sep_tr.findall(qn('w:tc'))
sep_data = ['机制变量', '', '', '']
for j, tc in enumerate(tcs):
    for p in tc.findall(qn('w:p')):
        runs = p.findall(qn('w:r'))
        if runs:
            t_el = runs[0].find(qn('w:t'))
            if t_el is not None:
                t_el.text = sep_data[j] if j < len(sep_data) else ''
            for r in runs[1:]:
                p.remove(r)

analyst_tr = t0.rows[14]._tr
t0._tbl.insert(list(t0._tbl).index(analyst_tr), sep_tr)
print("  Inserted '机制变量' separator row")

# Step 3: 在表格末尾添加 Disp 行
disp_tr = deepcopy(last_row._tr)
tcs = disp_tr.findall(qn('w:tc'))
disp_data = ['Disp', '分析师预测标准差/预测均值绝对值', '0.037', '0.046']
for j, tc in enumerate(tcs):
    for p in tc.findall(qn('w:p')):
        runs = p.findall(qn('w:r'))
        if runs:
            t_el = runs[0].find(qn('w:t'))
            if t_el is not None:
                t_el.text = disp_data[j] if j < len(disp_data) else ''
            for r in runs[1:]:
                p.remove(r)
t0._tbl.append(disp_tr)
print(f"  Added Disp row")

print("  Table 0 done!")

# ============================================================
# Save
# ============================================================
doc.save(doc_path)
print(f"\n=== Saved to {doc_path} ===")

# ============================================================
# Verify
# ============================================================
print("\n=== Verification ===")
doc2 = Document(doc_path)

print("\n--- Table 0 (last 7 rows) ---")
t0v = doc2.tables[0]
n = len(t0v.rows)
for i in range(n - 7, n):
    tcs = t0v.rows[i]._tr.findall(qn('w:tc'))
    cells = []
    for tc in tcs:
        ps = tc.findall(qn('w:p'))
        txt = ''
        for p in ps:
            for r in p.findall(qn('w:r')):
                t_el = r.find(qn('w:t'))
                if t_el is not None and t_el.text:
                    txt += t_el.text
        cells.append(txt.strip())
    print(f"  Row {i}: {cells}")

print("\n--- Table 5 (all rows via XML) ---")
t5v = doc2.tables[5]
for i in range(len(t5v.rows)):
    tcs = t5v.rows[i]._tr.findall(qn('w:tc'))
    cells = []
    for tc in tcs:
        ps = tc.findall(qn('w:p'))
        txt = ''
        for p in ps:
            for r in p.findall(qn('w:r')):
                t_el = r.find(qn('w:t'))
                if t_el is not None and t_el.text:
                    txt += t_el.text
        cells.append(txt.strip())
    print(f"  Row {i} ({len(tcs)} cells): {cells}")

print("\n--- Table 6 Panel B rows 14,16 ---")
t6v = doc2.tables[6]
for ri in [14, 16]:
    cells = [t6v.cell(ri, j).text.strip() for j in range(9)]
    print(f"  Row {ri}: {cells}")
