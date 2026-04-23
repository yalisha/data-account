"""
向v7.docx添加构念效度附表 + 修改para[56]和para[100]
"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import json, os

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC_PATH = f"{BASE}/manuscript/word/数据要素利用与资产定价效率v7.docx"

doc = Document(DOC_PATH)

with open(f"{BASE}/results/v11_expanded/construct_validity.json") as f:
    res = json.load(f)


def set_cell_text(cell, text, bold=False, size=Pt(9), align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = size
    run.font.name = '\u5b8b\u4f53'  # 宋体
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), '\u5b8b\u4f53')


def format_coef(coef, pval):
    stars = "***" if pval < 0.01 else "**" if pval < 0.05 else "*" if pval < 0.1 else ""
    return f"{coef:.4f}{stars}"


def format_t(t_val):
    return f"({t_val:.2f})"


# ============================================================
# 1. Modify para [56] - add CNRDS independence explanation
# ============================================================
new_56 = (
    "\u4e3a\u8fdb\u4e00\u6b65\u8bc4\u4f30\u6d4b\u5ea6\u7684\u5916\u90e8\u6548\u5ea6\uff0c"  # 为进一步评估测度的外部效度，
    "\u5c06\u672c\u6587\u7684DU_kw\u4e0e\u4e2d\u56fd\u7814\u7a76\u6570\u636e\u670d\u52a1\u5e73\u53f0\uff08CNRDS\uff09"  # 将本文的DU_kw与中国研究数据服务平台（CNRDS）
    "\u53d1\u5e03\u7684\u6570\u636e\u8981\u7d20\u6307\u6570\u8fdb\u884c\u5bf9\u7167\u3002"  # 发布的数据要素指数进行对照。
)

# Build para 56 text in plain form for readability
p56_text = (
    "为进一步评估测度的外部效度，"
    "将本文的DU_kw与中国研究数据服务平台（CNRDS）发布的数据要素指数进行对照。"
    "CNRDS数据要素指数由第三方学术机构独立编制，"
    "采用四维度关键词体系（数据要素存量、数据开发能力、数据驱动商业应用、数据价值变现），"
    "其关键词遴选标准和维度划分均独立于本文的五维度体系。"
    "两个指标虽然都基于年报文本，但关键词表的具体词条存在实质差异，"
    "因此相关性检验可以在词表构造层面提供收敛效度证据。"
    "在2018至2020年11,579个重叠的企业年度观测中，"
    "两个指标的Pearson相关系数为0.699、Spearman秩相关系数为0.787，"
    "均在1%水平上高度显著。较高的正相关表明DU_kw与外部权威指标方向一致，"
    "同时两者并非完全重合（r远低于1），反映了不同词表体系在测度口径上的合理差异。"
    "CNRDS数据要素指数仅覆盖三年，不足以支撑长面板的时序检验，"
    "但上述对照验证了本文指标在可比窗口内的收敛效度。"
)

p56 = doc.paragraphs[56]
for run in p56.runs:
    run.text = ""
if p56.runs:
    p56.runs[0].text = p56_text
else:
    p56.add_run(p56_text)

print("Para [56] updated.")

# ============================================================
# 2. Modify para [100] - add appendix table reference
# ============================================================
# Use the word "附表1" to reference the appendix table
p100_text = (
    "为排除DU_kw反映的是年报写作风格或泛化热词而非数据要素利用的可能，"
    "进行三组构念效度检验（详见附表1）。"
    "第一，在基准模型中分别加入年报总字数对数和MD\u0026A字数对数，"
    "控制\u201c话多\u201d和管理层叙事篇幅的影响。"
    "DU_kw系数在三种设定下均保持稳健（t值分别为-4.44、-4.11和-4.19），"
    "表明其解释力来自关键词的语义内容而非文本篇幅。"
    "第二，构造泛战略叙事词频（高质量发展、碳中和、乡村振兴等20个非数据相关的政策热词，"
    "每万字频率）作为安慰剂指标。泛战略词频单独回归时边际显著（t=-2.10），"
    "但与DU_kw同时纳入后变为不显著（t=-1.18，p=0.239），"
    "而DU_kw系数几乎不变（t=-4.19）。"
    "进一步从关键词表中剔除数字化转型、人工智能、大数据等8个最泛化的词后，"
    "重新计算的DU_kw_strict仍然显著（t=-3.90，p<0.001）。"
    "上述检验表明，DU_kw捕捉的是数据要素利用的特定语义信号，"
    "而非管理层赶热点式的泛化表述。"
)

p100 = doc.paragraphs[100]
for run in p100.runs:
    run.text = ""
if p100.runs:
    p100.runs[0].text = p100_text
else:
    p100.add_run(p100_text)

print("Para [100] updated.")

# ============================================================
# 3. Add appendix table after references
# ============================================================

# Heading
doc.add_paragraph("")
heading_p = doc.add_paragraph("")
heading_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = heading_p.add_run("\u9644\u88681  \u6784\u5ff5\u6548\u5ea6\u68c0\u9a8c")  # 附表1  构念效度检验
run.bold = True
run.font.size = Pt(12)
run.font.name = '\u5b8b\u4f53'

# Panel A label
panel_a_p = doc.add_paragraph("")
panel_a_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = panel_a_p.add_run("Panel A: \u533a\u5206\u6027\u4e0e\u5b89\u6170\u5242\u68c0\u9a8c\uff08\u56e0\u53d8\u91cf = PriceDelay\uff09")
run.bold = True
run.font.size = Pt(10)

# Panel A table: 15 rows x 7 cols
n_rows = 15
n_cols = 7
table_a = doc.add_table(rows=n_rows, cols=n_cols)
table_a.alignment = WD_TABLE_ALIGNMENT.CENTER

# Header
headers = ["", "(1)", "(2)", "(3)", "(4)", "(5)", "(6)"]
for j, h in enumerate(headers):
    set_cell_text(table_a.cell(0, j), h, bold=True)

# DU_kw (rows 1-2)
set_cell_text(table_a.cell(1, 0), "DU_kw", align=WD_ALIGN_PARAGRAPH.LEFT)
# col 1: horserace_total_chars
set_cell_text(table_a.cell(1, 1), format_coef(res['horserace_total_chars']['DU_kw_coef'], res['horserace_total_chars']['DU_kw_p']))
set_cell_text(table_a.cell(2, 1), format_t(res['horserace_total_chars']['DU_kw_t']))
# col 2: horserace_mda_chars
set_cell_text(table_a.cell(1, 2), format_coef(res['horserace_mda_chars']['DU_kw_coef'], res['horserace_mda_chars']['DU_kw_p']))
set_cell_text(table_a.cell(2, 2), format_t(res['horserace_mda_chars']['DU_kw_t']))
# col 3: horserace_both
set_cell_text(table_a.cell(1, 3), format_coef(res['horserace_both']['DU_kw_coef'], res['horserace_both']['DU_kw_p']))
set_cell_text(table_a.cell(2, 3), format_t(res['horserace_both']['DU_kw_t']))
# col 4: placebo only - no DU_kw
# col 5: horserace_placebo
set_cell_text(table_a.cell(1, 5), format_coef(res['horserace_placebo']['DU_kw_coef'], res['horserace_placebo']['DU_kw_p']))
set_cell_text(table_a.cell(2, 5), format_t(res['horserace_placebo']['DU_kw_t']))
# col 6: DU_kw_strict - uses different variable

# ln_total_chars (rows 3-4)
set_cell_text(table_a.cell(3, 0), "ln_total_chars", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_a.cell(3, 1), format_coef(res['horserace_total_chars']['ln_total_chars_coef'], 0.008))
set_cell_text(table_a.cell(4, 1), format_t(res['horserace_total_chars']['ln_total_chars_t']))
set_cell_text(table_a.cell(3, 3), "\u63a7\u5236")  # 控制

# ln_mda_chars (rows 5-6)
set_cell_text(table_a.cell(5, 0), "ln_mda_chars", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_a.cell(5, 2), format_coef(res['horserace_mda_chars']['ln_mda_chars_coef'], 0.135))
set_cell_text(table_a.cell(6, 2), format_t(res['horserace_mda_chars']['ln_mda_chars_t']))
set_cell_text(table_a.cell(5, 3), "\u63a7\u5236")

# Placebo_kw (rows 7-8)
set_cell_text(table_a.cell(7, 0), "Placebo_kw", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_a.cell(7, 4), format_coef(res['placebo_only']['Placebo_kw_coef'], res['placebo_only']['Placebo_kw_p']))
set_cell_text(table_a.cell(8, 4), format_t(res['placebo_only']['Placebo_kw_t']))
set_cell_text(table_a.cell(7, 5), format_coef(res['horserace_placebo']['Placebo_kw_coef'], res['horserace_placebo']['Placebo_kw_p']))
set_cell_text(table_a.cell(8, 5), format_t(res['horserace_placebo']['Placebo_kw_t']))

# DU_kw_strict (rows 9-10)
set_cell_text(table_a.cell(9, 0), "DU_kw_strict", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_a.cell(9, 6), format_coef(res['du_kw_strict']['coef'], res['du_kw_strict']['p']))
set_cell_text(table_a.cell(10, 6), format_t(res['du_kw_strict']['t']))

# Controls / FE / N (rows 11-14)
set_cell_text(table_a.cell(11, 0), "\u63a7\u5236\u53d8\u91cf", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_a.cell(12, 0), "\u4f01\u4e1a\u56fa\u5b9a\u6548\u5e94", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_a.cell(13, 0), "\u5e74\u4efd\u56fa\u5b9a\u6548\u5e94", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_a.cell(14, 0), "N", align=WD_ALIGN_PARAGRAPH.LEFT)

for j in range(1, 7):
    set_cell_text(table_a.cell(11, j), "\u662f")  # 是
    set_cell_text(table_a.cell(12, j), "\u662f")
    set_cell_text(table_a.cell(13, j), "\u662f")

Ns = [
    res['horserace_total_chars']['N'],
    res['horserace_mda_chars']['N'],
    res['horserace_both']['N'],
    res['placebo_only']['N'],
    res['horserace_placebo']['N'],
    res['du_kw_strict']['N'],
]
for j, n in enumerate(Ns):
    set_cell_text(table_a.cell(14, j+1), f"{n:,}")

# Panel B
panel_b_p = doc.add_paragraph("")
run = panel_b_p.add_run(
    "Panel B: \u6536\u655b\u6548\u5ea6\u68c0\u9a8c\uff08DU_kw\u4e0eCNRDS\u6570\u636e\u8981\u7d20\u6307\u6570\uff0c2018-2020\uff09"
)
run.bold = True
run.font.size = Pt(10)

table_b = doc.add_table(rows=4, cols=3)
table_b.alignment = WD_TABLE_ALIGNMENT.CENTER

set_cell_text(table_b.cell(0, 0), "\u76f8\u5173\u6027\u6307\u6807", bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_b.cell(0, 1), "\u7cfb\u6570", bold=True)
set_cell_text(table_b.cell(0, 2), "p\u503c", bold=True)

cnrds = res['cnrds_validity']
set_cell_text(table_b.cell(1, 0), "Pearson\u76f8\u5173\u7cfb\u6570", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_b.cell(1, 1), f"{cnrds['pearson_r']:.4f}***")
set_cell_text(table_b.cell(1, 2), "<0.001")

set_cell_text(table_b.cell(2, 0), "Spearman\u79e9\u76f8\u5173\u7cfb\u6570", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_b.cell(2, 1), f"{cnrds['spearman_r']:.4f}***")
set_cell_text(table_b.cell(2, 2), "<0.001")

set_cell_text(table_b.cell(3, 0), "\u539f\u59cb\u8ba1\u6570Pearson\u76f8\u5173\u7cfb\u6570", align=WD_ALIGN_PARAGRAPH.LEFT)
set_cell_text(table_b.cell(3, 1), f"{cnrds['raw_count_pearson']:.4f}***")
set_cell_text(table_b.cell(3, 2), "<0.001")

# Table note
note_p = doc.add_paragraph("")
note_text = (
    "\u6ce8\uff1aPanel A\u56e0\u53d8\u91cf\u4e3aPriceDelay\uff0c"
    "\u6240\u6709\u6a21\u578b\u5305\u542b\u4f01\u4e1a\u548c\u5e74\u4efd\u56fa\u5b9a\u6548\u5e94\uff0c"
    "\u6807\u51c6\u8bef\u6309\u884c\u4e1a\u00d7\u5e74\u4efd\u805a\u7c7b\u3002\u62ec\u53f7\u5185\u4e3at\u503c\u3002"
    "\u5217(1)-(3)\u4e3a\u6587\u672c\u7bc7\u5e45\u63a7\u5236\u7684horse-race\u68c0\u9a8c\uff0c"
    "\u5217(1)\u52a0\u5165\u5e74\u62a5\u603b\u5b57\u6570\u5bf9\u6570\uff0c"
    "\u5217(2)\u52a0\u5165MD&A\u5b57\u6570\u5bf9\u6570\uff0c"
    "\u5217(3)\u540c\u65f6\u52a0\u5165\u4e24\u8005\u3002"
    "\u5217(4)-(5)\u4e3a\u6cdb\u6218\u7565\u70ed\u8bcd\u5b89\u6170\u5242\u68c0\u9a8c\uff0c"
    "\u5217(4)\u4ec5\u7eb3\u5165Placebo_kw\uff0c\u5217(5)\u540c\u65f6\u7eb3\u5165DU_kw\u548cPlacebo_kw\u3002"
    "\u5217(6)\u4e3a\u5254\u96648\u4e2a\u6cdb\u5316\u5173\u952e\u8bcd\u540e\u7684DU_kw_strict\u68c0\u9a8c\u3002"
    "Panel B\u4e2dN=11,579\uff0c\u4e3a2018-2020\u5e74DU_kw\u4e0eCNRDS\u6570\u636e\u8981\u7d20\u6307\u6570\u7684\u91cd\u53e0\u6837\u672c\u3002"
    "CNRDS\u6570\u636e\u8981\u7d20\u6307\u6570\u7531\u7b2c\u4e09\u65b9\u5b66\u672f\u673a\u6784\u72ec\u7acb\u7f16\u5236\uff0c"
    "\u91c7\u7528\u56db\u7ef4\u5ea6\u5173\u952e\u8bcd\u4f53\u7cfb\u3002"
    "***\u3001**\u3001*\u5206\u522b\u8868\u793a\u57281%\u30015%\u300110%\u6c34\u5e73\u4e0a\u663e\u8457\u3002"
)
run = note_p.add_run(note_text)
run.font.size = Pt(8)
run.font.name = '\u5b8b\u4f53'

doc.save(DOC_PATH)
print("\nDone! Saved v7.docx with:")
print("  - Para [56]: CNRDS independence explanation added")
print("  - Para [100]: Appendix table reference added")
print("  - Appendix Table 1: Construct validity (Panel A + Panel B)")
