"""
v8格式整理：
1. Para [102] 恢复标题3样式（被replace_para改成Normal了）
2. 清理所有段落中的空runs（replace_para留下的残余）
3. 修复双空格
"""
from docx import Document
from lxml import etree
import os
import copy

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC_PATH = f"{BASE}/manuscript/word/数据要素利用与资产定价效率v8.docx"
doc = Document(DOC_PATH)

ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# ============================================================
# 1. 修复 Para [102] 标题样式
# ============================================================
p102 = doc.paragraphs[102]
print(f"[1] Para [102] style: {p102.style.name} -> 标题3")
p102.style = doc.styles['标题3']

# ============================================================
# 2. 清理空runs
# ============================================================
empty_cleaned = 0
for i, p in enumerate(doc.paragraphs):
    runs_to_remove = []
    has_content_run = False
    for r in p.runs:
        if r.text == '' or r.text is None:
            runs_to_remove.append(r)
        else:
            has_content_run = True

    # Only remove empty runs if there's at least one content run
    if has_content_run and runs_to_remove:
        for r in runs_to_remove:
            r._element.getparent().remove(r._element)
            empty_cleaned += 1

print(f"[2] Cleaned {empty_cleaned} empty runs across all paragraphs")

# ============================================================
# 3. 修复双空格
# ============================================================
double_space_fixed = 0
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        if '  ' in r.text:
            r.text = r.text.replace('  ', ' ')
            double_space_fixed += 1
            print(f"  Fixed double space in Para [{i}]")

print(f"[3] Fixed {double_space_fixed} double-space instances")

# Save
doc.save(DOC_PATH)
print(f"\nDone! v8.docx formatting cleaned.")
