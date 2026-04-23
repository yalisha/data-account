"""
投资组合检验降格：从独立视角改为补充分析
[1] Para [102]: 节标题加"补充分析"
[2] Para [103]: 引入语从"转向资产定价视角"改为补充性定位
[3] Para [108]: 末句强调描述性质
"""
from docx import Document
import os

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC_PATH = f"{BASE}/manuscript/word/数据要素利用与资产定价效率v7.docx"
doc = Document(DOC_PATH)


def replace_para(idx, new_text):
    p = doc.paragraphs[idx]
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = new_text
    else:
        p.add_run(new_text)
    print(f"  Para [{idx}] updated ({len(new_text)} chars)")


# ============================================================
# Para [102]: 节标题
# ============================================================
new_102 = "（四）补充分析：投资组合检验"
print("Updating Para [102] (section title)...")
replace_para(102, new_102)


# ============================================================
# Para [103]: 引入语
# ============================================================
new_103 = (
    "面板回归从微观层面确认了数据要素利用降低股价延迟的效应。"
    "作为补充，以下用投资组合方法考察数据利用差异是否在截面收益率上留下痕迹。"
)
print("Updating Para [103] (portfolio intro)...")
replace_para(103, new_103)


# ============================================================
# Para [108]: DAT因子段落 — 末句明确描述性质
# ============================================================
new_108 = (
    "基于Q5-Q1构造的数据资产因子（DAT）在FF3模型下的月度收益率为0.37%"
    "（t=1.17），未达到统计显著水平。"
    "这与非对称定价的模式吻合："
    "市场对低数据利用企业的折价远大于对高利用企业的溢价。"
    "数据利用的改善在于消解信息模糊带来的折价，"
    "而不是创造正向超额收益。"
    "低数据利用企业面临更大的信息融入障碍，"
    "市场对它们的定价偏离更为明显。"
    "需要说明的是，上述结果是描述性的，"
    "并不意味着数据利用构成一个独立的定价因子。"
)
print("Updating Para [108] (DAT factor, add descriptive caveat)...")
replace_para(108, new_108)


# Save
doc.save(DOC_PATH)
print(f"\nDone! 3 paragraphs updated in v7.docx")
