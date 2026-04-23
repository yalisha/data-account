"""
修复v9.docx两个问题：
1. [103] 梯度检验 p=0.141 措辞改为审慎表述
2. [115][117][121][123] 异质性分析各小节补充参考文献
"""

from docx import Document

doc_path = "/Users/mac/computerscience/15会计研究/manuscript/word/数据要素利用与资产定价效率v9.docx"
doc = Document(doc_path)


def replace_in_para(idx, old, new):
    p = doc.paragraphs[idx]
    if old not in p.text:
        print(f"  [WARNING] [{idx}] NOT FOUND: '{old}'")
        return False
    for run in p.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            print(f"  [{idx}] replaced OK")
            return True
    # concat fallback
    combined = ''.join(r.text for r in p.runs)
    if old in combined:
        new_text = combined.replace(old, new)
        for r in p.runs[1:]:
            p._element.remove(r._element)
        p.runs[0].text = new_text
        print(f"  [{idx}] concat-replaced OK")
        return True
    print(f"  [WARNING] [{idx}] run fallback failed")
    return False


# ============================================================
# Fix 1: [103] 梯度检验措辞
# p=0.141不显著，"进一步验证"太强
# ============================================================
print("=== Fix 1: [103] 梯度检验措辞 ===")
replace_in_para(
    103,
    "这一梯度模式进一步验证了信息中介在数据信息传导中的作用",
    "这一梯度模式与信息中介渠道的理论预期方向一致，但未达到传统显著水平"
)

# ============================================================
# Fix 2: [115] 产权性质 - 补Diamond和Verrecchia(1991)
# ============================================================
print("\n=== Fix 2: [115] 产权性质补文献 ===")
replace_in_para(
    115,
    "在信息透明度较低的基线条件下，数据要素利用所释放的增量信息具有更高的边际价值",
    "在信息透明度较低的基线条件下，增量披露对降低信息不对称的边际效果更强（Diamond和Verrecchia，1991），数据要素利用所释放的增量信息因此具有更高的边际价值"
)

# ============================================================
# Fix 3: [117] 流动性 - 补Chordia等(2008)
# ============================================================
print("\n=== Fix 3: [117] 流动性补文献 ===")
replace_in_para(
    117,
    "数据要素利用所释放的增量信息可能无法充分定价",
    "数据要素利用所释放的增量信息可能无法充分定价。Chordia等（2008）发现流动性改善能显著提升市场效率，这意味着数据利用的定价效率效应在流动性充裕的市场中传导更为顺畅"
)

# ============================================================
# Fix 4: [121] 规模 - 补Hou和Moskowitz(2005)、Hong等(2000)
# ============================================================
print("\n=== Fix 4: [121] 规模补文献 ===")
replace_in_para(
    121,
    "大企业的信息环境更为透明、分析师覆盖更广，数据信息的边际增量价值可能较低",
    "Hou和Moskowitz（2005）发现小企业的股价延迟系统性高于大企业，Hong等（2000）进一步指出这与分析师覆盖不足导致的信息扩散缓慢有关，意味着数据信息对小企业的边际增量价值可能更高"
)

# ============================================================
# Fix 5: [123] 行业 - 补Dessaint等(2024)
# ============================================================
print("\n=== Fix 5: [123] 行业补文献 ===")
# Check if Dessaint is in the bibliography
has_dessaint = any("Dessaint" in p.text for p in doc.paragraphs[130:170])
print(f"  Dessaint in bibliography: {has_dessaint}")

replace_in_para(
    123,
    "高科技行业由于技术导向强、信息更新频繁、市场关注度高，信息环境的基线透明度相对较高",
    "高科技行业由于技术导向强、信息更新频繁、市场关注度高，信息环境的基线透明度相对较高（Dessaint等，2024）"
)

# ============================================================
# Save
# ============================================================
doc.save(doc_path)
print(f"\n=== Saved ===")

# ============================================================
# Verify
# ============================================================
print("\n=== Verification ===")
doc2 = Document(doc_path)

# [103] check
t103 = doc2.paragraphs[103].text
idx = t103.find("梯度模式")
if idx >= 0:
    print(f"\n[103] 梯度: ...{t103[idx:idx+60]}...")
else:
    print(f"\n[103] WARNING: '梯度模式' not found")

# [115] check
t115 = doc2.paragraphs[115].text
print(f"\n[115] Diamond: {'Diamond' in t115}")
print(f"  snippet: ...{t115[t115.find('Diamond')-20:t115.find('Diamond')+80] if 'Diamond' in t115 else 'NOT FOUND'}...")

# [117] check
t117 = doc2.paragraphs[117].text
print(f"\n[117] Chordia: {'Chordia' in t117}")

# [121] check
t121 = doc2.paragraphs[121].text
print(f"\n[121] Hou: {'Hou' in t121}")
print(f"  Hong: {'Hong' in t121}")

# [123] check
t123 = doc2.paragraphs[123].text
print(f"\n[123] Dessaint: {'Dessaint' in t123}")
