"""
理论分析章节重构：
1. 子标题去AI味：（二）改名，（三）删除
2. 机制渠道精简为概念铺垫（删实证文献引用，留给机制分析节）
3. 异质性4段合并为1段桥接
处理顺序：从后往前（先删段落，再改文本）
"""

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy
import os

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOC_PATH = f"{BASE}/manuscript/word/数据要素利用与资产定价效率v8.docx"

doc = Document(DOC_PATH)


def replace_para_text(para, new_text):
    """Clear all runs, add single new run with first run's formatting."""
    runs = para.runs
    first_rPr = None
    if runs:
        rPr_elem = runs[0]._element.find(qn('w:rPr'))
        if rPr_elem is not None:
            first_rPr = copy.deepcopy(rPr_elem)

    for r in runs:
        r._element.getparent().remove(r._element)
    for r_elem in list(para._element.findall(qn('w:r'))):
        para._element.remove(r_elem)

    new_run = OxmlElement('w:r')
    if first_rPr is not None:
        new_run.append(first_rPr)
    t_elem = OxmlElement('w:t')
    t_elem.text = new_text
    t_elem.set(qn('xml:space'), 'preserve')
    new_run.append(t_elem)
    para._element.append(new_run)


def delete_para(doc, idx):
    """Delete paragraph at index."""
    p = doc.paragraphs[idx]
    p._element.getparent().remove(p._element)


# ============================================================
# Verify current structure
# ============================================================
print("=== 当前结构 ===")
for i in range(25, 44):
    p = doc.paragraphs[i]
    style = p.style.name
    text = p.text[:80] if p.text else "(empty)"
    print(f"[{i}] {style}: {text}")
print()

# Assertions
assert '数据要素利用与资本市场定价效率' in doc.paragraphs[26].text
assert '传导机制' in doc.paragraphs[31].text or '信息中介' in doc.paragraphs[31].text
assert '边界条件' in doc.paragraphs[39].text or '异质性' in doc.paragraphs[39].text


# ============================================================
# Step 1: 从后往前处理异质性部分 [39]-[43]
# 删除[40]-[43]，用[40]替换为桥接段，删除[39]标题
# ============================================================

print("=== Step 1: 异质性部分 ===")

# 桥接段内容
het_bridge = (
    "上述效应的强度可能因企业特征和市场条件而异。"
    "后续分析从产权性质、市场流动性和行业属性三个维度检验效应的边界条件。"
    "产权性质反映信息环境的制度基础，市场流动性反映信息融入价格的物理约束，"
    "行业属性反映信息环境的基线透明度差异。"
)

# 从后往前删除 [43], [42], [41]
for i in [43, 42, 41]:
    print(f"  Deleting [{i}]: {doc.paragraphs[i].text[:50]}...")
    delete_para(doc, i)

# Now [40] is the old [40], replace it with bridge
p40 = doc.paragraphs[40]
print(f"  [{40}] was: {p40.text[:50]}...")
replace_para_text(p40, het_bridge)
print(f"  [{40}] now: {het_bridge[:60]}...")

# Delete [39] heading (边界条件)
p39 = doc.paragraphs[39]
print(f"  Deleting [{39}] heading: {p39.text}")
delete_para(doc, 39)

print()


# ============================================================
# Step 2: 精简机制渠道 [33]-[38]
# 注意：经Step1删除4段后，索引已经偏移！但33-38在39之前，未受影响
# 实际上我们先处理了39+之后的段落，33-38仍在原位
# 等等——delete_para用的是当时的索引。Step1删了[43][42][41][39]（都在39+）
# 所以[33]-[38]未受影响
# ============================================================

print("=== Step 2: 精简机制渠道 ===")

# [31] 改标题
p31 = doc.paragraphs[31]
print(f"  [{31}] was: {p31.text}")
replace_para_text(p31, "（二）传导机制与研究假说")
print(f"  [{31}] now: （二）传导机制与研究假说")

# [33]-[34] 合并为1段概念性论述（分析师渠道）
analyst_concept = (
    "第一条渠道是信息中介渠道。"
    "分析师连接着企业信息披露与市场定价，其覆盖密度直接影响信息向价格的转换效率。"
    "数据要素利用涉及数据资产化、数字化转型等技术含量较高的业务实践，"
    "年报中此类信息的增加为分析师提供了新的研究素材和增量信息来源，"
    "可能吸引更多分析师关注和跟踪，从而加速企业特质信息向价格的传导。据此提出："
)

p33 = doc.paragraphs[33]
print(f"  [{33}] was: {p33.text[:60]}...")
replace_para_text(p33, analyst_concept)
print(f"  [{33}] now: {analyst_concept[:60]}...")

# 删除[34]（分析师渠道续段，已合并到[33]）
p34 = doc.paragraphs[34]
print(f"  Deleting [{34}]: {p34.text[:60]}...")
delete_para(doc, 34)

# 现在原[35]变成[34]（H2a声明），原[36]变成[35]...
# [35] = 原[36] Disp渠道 → 精简
disp_concept = (
    "第二条渠道是信息质量渠道。"
    "数据驱动的业务模式产生了更多结构化的运营信息，"
    "使分析师在预测企业未来盈利时面临的不确定性降低，从而缩小预测分歧（Zhang，2006）。"
    "预测分歧度的下降意味着市场参与者对企业价值的判断趋于一致，"
    "信息不对称程度减弱，价格发现过程得以加速。据此提出："
)

# 找到当前的Disp渠道段（原[36]，删除[34]后变成[35]）
p_disp = doc.paragraphs[35]
print(f"  [{35}] was: {p_disp.text[:60]}...")
replace_para_text(p_disp, disp_concept)
print(f"  [{35}] now: {disp_concept[:60]}...")

# [37] = 原[38] Amihud渠道（原[38]经过删除[34]后变成[37]）→ 大幅精简
amihud_concept = (
    "第三条渠道是市场流动性渠道。"
    "信息不对称加剧知情与非知情交易者之间的逆向选择，"
    "扩大买卖价差，降低市场流动性，阻碍信息向价格的传递。"
    "数据要素利用通过增加企业可观测信息的供给、缩小信息不对称程度，"
    "有望缓解逆向选择问题，改善市场流动性。"
    "与前两条渠道侧重信息的生产和加工不同，"
    "市场流动性渠道刻画的是信息融入价格的交易环节，"
    "三者共同构成了从信息供给到定价效率改善的完整链条。"
    "本文在机制分析中以Amihud非流动性指标作为中间变量加以检验。"
)

p_amihud = doc.paragraphs[37]
print(f"  [{37}] was: {p_amihud.text[:60]}...")
replace_para_text(p_amihud, amihud_concept)
print(f"  [{37}] now: {amihud_concept[:60]}...")

print()


# ============================================================
# Step 3: 验证最终结构
# ============================================================
print("=== 验证最终结构 ===")
for i in range(25, 45):
    if i >= len(doc.paragraphs):
        break
    p = doc.paragraphs[i]
    style = p.style.name
    text = p.text[:90] if p.text else "(empty)"
    print(f"[{i}] {style}: {text}")

# 检查关键点
found_old_heading = False
for i in range(25, 45):
    if i >= len(doc.paragraphs):
        break
    p = doc.paragraphs[i]
    if '边界条件' in p.text:
        found_old_heading = True
        print(f"\nWARNING: '边界条件' still found at [{i}]!")

if not found_old_heading:
    print("\n✓ '边界条件' heading removed successfully")

doc.save(DOC_PATH)
print(f"\nDone! Saved to {DOC_PATH}")
