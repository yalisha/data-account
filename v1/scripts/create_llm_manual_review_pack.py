"""
Create a manual review pack for the 120-context LLM validation sample.

Outputs:
  - results/v15_llm_review/llm_manual_review_sheet.xlsx
  - results/v15_llm_review/llm_manual_review_sheet.csv
  - results/v15_llm_review/llm_manual_review_guide.md
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


BASE = Path(__file__).resolve().parents[1]
SAMPLE_CSV = BASE / "results" / "v14" / "llm_validation_sample.csv"
OUT_DIR = BASE / "results" / "v15_llm_review"
OUT_DIR.mkdir(parents=True, exist_ok=True)
CODEX_SCORE_CSV = OUT_DIR / "codex_independent_rerun_scores.csv"

OUT_XLSX = OUT_DIR / "llm_manual_review_sheet.xlsx"
OUT_CSV = OUT_DIR / "llm_manual_review_sheet.csv"
OUT_GUIDE = OUT_DIR / "llm_manual_review_guide.md"


RUBRIC_ROWS = [
    (
        0,
        "仅为口号式、背景式或无关提及",
        "没有企业层面的具体数据对象、动作、场景或治理含义",
        "税收优惠列表、子公司名称、政策口号、行业背景、泛泛而谈的数字化表述",
    ),
    (
        1,
        "提到建设方向或能力，但缺乏清晰落地场景",
        "出现平台、信息化建设、能力提升、推进数字化等，但缺少具体业务对象和动作",
        "“推进信息化建设”“加强数据能力”“搭建平台”之类方向性描述",
    ),
    (
        2,
        "出现明确的数据应用场景、治理动作或决策用途",
        "可识别具体对象、业务环节、管理动作或应用目的",
        "客户数据库、需求预测、风控模型、库存调度、数据治理动作",
    ),
    (
        3,
        "在2分基础上进一步出现系统嵌入、流程闭环、量化结果或价值实现",
        "同时出现平台/模型/流程 + 具体用途 + 明确实施方式、结果或价值实现",
        "工业互联网平台+采集传输+模型+检修策略；数据资产化或价值实现有清晰证据",
    ),
]


DECISION_RULES = [
    "只基于该条关键词上下文打分，不依据你对企业的外部了解补充信息。",
    "判断核心是“数据资源是否被实际用于经营、治理或价值实现”，不是“是否提到数字化”。",
    "若样本同时包含泛化叙事和具体应用，以更能代表该窗口主旨的内容为准。",
    "若在相邻两档之间难以判断，默认取较低分，避免过度打高分。",
    "若上下文主要是名单、主体名称、资质备案、政策复述、宏观口号，优先判为0分。",
    "若你认为该样本存在歧义或上下文截断影响判断，请在备注列写明原因。",
]


def build_review_frame(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out = out.rename(columns={"parent_llm_score": "archived_parent_score"})
    if CODEX_SCORE_CSV.exists():
        codex = pd.read_csv(CODEX_SCORE_CSV)
        out = out.merge(codex, on=["sample_id", "keyword"], how="left", validate="1:1")
    else:
        out["codex_independent_score"] = ""
        out["codex_confidence"] = ""
        out["codex_notes"] = ""
    out["human_score"] = ""
    out["human_confidence"] = ""
    out["human_notes"] = ""
    out["score_gap_vs_parent"] = ""
    out["needs_adjudication"] = ""
    cols = [
        "sample_id",
        "Stkcd",
        "year",
        "keyword",
        "context",
        "generic_bucket",
        "score_bucket",
        "density_bucket",
        "window_hits",
        "archived_parent_score",
        "codex_independent_score",
        "codex_confidence",
        "codex_notes",
        "human_score",
        "human_confidence",
        "score_gap_vs_parent",
        "needs_adjudication",
        "human_notes",
    ]
    return out[cols]


def write_guide(df: pd.DataFrame) -> None:
    strat_counts = (
        df.groupby(["generic_bucket", "score_bucket", "density_bucket"])
        .size()
        .reset_index(name="n")
        .sort_values(["generic_bucket", "score_bucket", "density_bucket"])
    )
    lines = [
        "# LLM 人工复核说明",
        "",
        "## 复核目标",
        "",
        "- 对 120 条上下文样本进行独立人工打分，检验归档语义评分与人工判断的一致性。",
        "- 评分对象是关键词命中的局部上下文，而不是企业整体数字化程度。",
        "",
        "## 判分标准",
        "",
    ]
    for score, label, cue, example in RUBRIC_ROWS:
        lines.extend(
            [
                f"### {score} 分",
                f"- 定义：{label}",
                f"- 主要判断信号：{cue}",
                f"- 常见场景：{example}",
                "",
            ]
        )
    lines.extend(
        [
            "## 决策规则",
            "",
        ]
    )
    lines.extend([f"- {rule}" for rule in DECISION_RULES])
    lines.extend(
        [
            "",
            "## 审阅文件字段",
            "",
            "- `archived_parent_score`：当前归档的企业-年份层面众数分值，仅供对照。",
            "- `codex_independent_score`：独立 Codex worker 基于同一 rubric 给出的样本分值，可与人工评分对比。",
            "- `codex_confidence` / `codex_notes`：独立评分的置信度与简短备注。",
            "- `human_score`：请填写你的人工评分（0/1/2/3）。",
            "- `human_confidence`：建议填写 `high / medium / low`。",
            "- `score_gap_vs_parent`：可在人工评分后填写与归档分值的差值。",
            "- `needs_adjudication`：若你认为需要二次讨论，可填 `Y`。",
            "- `human_notes`：记录歧义点、上下文截断、或为何与你的直觉分值不同。",
            "",
            "## 分层分布",
            "",
        ]
    )
    for row in strat_counts.itertuples(index=False):
        lines.append(
            f"- {row.generic_bucket} / {row.score_bucket} / {row.density_bucket}: {row.n}"
        )

    OUT_GUIDE.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_workbook(df: pd.DataFrame) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "samples"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    sub_fill = PatternFill("solid", fgColor="F4F6F8")
    wrap = Alignment(vertical="top", wrap_text=True)
    top = Alignment(vertical="top")
    bold = Font(bold=True)

    for row_idx, row in enumerate([list(df.columns)] + df.astype(str).values.tolist(), start=1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = wrap if col_idx in {5, 13, 18} else top
            if row_idx == 1:
                cell.font = bold
                cell.fill = header_fill

    ws.freeze_panes = "A2"
    widths = {
        "A": 11,
        "B": 10,
        "C": 8,
        "D": 12,
        "E": 95,
        "F": 14,
        "G": 12,
        "H": 13,
        "I": 10,
        "J": 18,
        "K": 12,
        "L": 16,
        "M": 24,
        "N": 12,
        "O": 16,
        "P": 16,
        "Q": 16,
        "R": 36,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    for row in range(2, len(df) + 2):
        ws.row_dimensions[row].height = 72

    score_dv = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    confidence_dv = DataValidation(type="list", formula1='"high,medium,low"', allow_blank=True)
    adjudication_dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(score_dv)
    ws.add_data_validation(confidence_dv)
    ws.add_data_validation(adjudication_dv)
    score_dv.add(f"N2:N{len(df)+1}")
    confidence_dv.add(f"O2:O{len(df)+1}")
    adjudication_dv.add(f"Q2:Q{len(df)+1}")

    rubric_ws = wb.create_sheet("rubric")
    rubric_ws.append(["score", "定义", "主要判断信号", "常见场景"])
    for row in RUBRIC_ROWS:
        rubric_ws.append(list(row))
    for cell in rubric_ws[1]:
        cell.font = bold
        cell.fill = header_fill
    for row in rubric_ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    rubric_ws.column_dimensions["A"].width = 8
    rubric_ws.column_dimensions["B"].width = 28
    rubric_ws.column_dimensions["C"].width = 40
    rubric_ws.column_dimensions["D"].width = 52

    inst_ws = wb.create_sheet("instructions")
    inst_ws["A1"] = "人工复核说明"
    inst_ws["A1"].font = Font(bold=True, size=14)
    inst_ws["A3"] = "复核目标"
    inst_ws["A3"].font = bold
    inst_ws["A4"] = "对 120 条上下文样本做独立人工打分，检验归档语义评分与人工判断的一致性。"
    inst_ws["A6"] = "决策规则"
    inst_ws["A6"].font = bold
    for i, rule in enumerate(DECISION_RULES, start=7):
        inst_ws[f"A{i}"] = f"{i-6}. {rule}"
    inst_ws["A15"] = "填写要求"
    inst_ws["A15"].font = bold
    inst_ws["A16"] = "请在 samples 工作表填写 human_score、human_confidence、human_notes；若需要二次讨论，在 needs_adjudication 填 Y。"
    inst_ws["A18"] = "字段提示"
    inst_ws["A18"].font = bold
    field_notes = [
        "archived_parent_score：当前归档的企业-年份层面众数分值，仅供对照。",
        "codex_independent_score：独立 Codex worker 的样本级评分，可作参考对照。",
        "codex_confidence / codex_notes：独立评分置信度及简短备注。",
        "human_score：你的人工评分（0/1/2/3）。",
        "human_confidence：建议填写 high / medium / low。",
        "score_gap_vs_parent：可在人工评分后填写与归档分值的差值。",
        "needs_adjudication：若需要二次讨论，填 Y。",
        "human_notes：记录歧义点、上下文截断或打分理由。",
    ]
    for i, note in enumerate(field_notes, start=19):
        inst_ws[f"A{i}"] = f"- {note}"

    for ws2 in [rubric_ws, inst_ws]:
        for row in ws2.iter_rows():
            for cell in row:
                cell.alignment = wrap
    inst_ws.column_dimensions["A"].width = 120

    stats_ws = wb.create_sheet("strata")
    stats_ws.append(["generic_bucket", "score_bucket", "density_bucket", "n"])
    counts = (
        df.groupby(["generic_bucket", "score_bucket", "density_bucket"])
        .size()
        .reset_index(name="n")
    )
    for row in counts.itertuples(index=False):
        stats_ws.append(list(row))
    for cell in stats_ws[1]:
        cell.font = bold
        cell.fill = sub_fill
    for col in ["A", "B", "C", "D"]:
        stats_ws.column_dimensions[col].width = 18

    wb.save(OUT_XLSX)


def main() -> None:
    df = pd.read_csv(SAMPLE_CSV)
    review_df = build_review_frame(df)
    review_df.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")
    write_guide(review_df)
    write_workbook(review_df)
    print(f"Saved review workbook to {OUT_XLSX}")
    print(f"Saved review csv to {OUT_CSV}")
    print(f"Saved guide to {OUT_GUIDE}")


if __name__ == "__main__":
    main()
