"""
Generate manuscript v11 Word document from markdown chapters.
Features: OMML equations, three-line tables (三线表), subtitle, equation numbering.
"""
import re, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, Emu, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

BASE = Path("/Users/mac/computerscience/15会计研究")
OUTPUT = BASE / "manuscript/word/数据要素利用与资产定价效率v11.docx"

MD_FILES = [
    BASE / "manuscript/00_摘要.md",
    BASE / "manuscript/01_引言.md",
    BASE / "manuscript/02_理论与假说.md",
    BASE / "manuscript/03_研究设计.md",
    BASE / "manuscript/04_实证检验.md",
    BASE / "manuscript/05_进一步分析.md",
    BASE / "manuscript/06_异质性分析.md",
    BASE / "manuscript/07_结论.md",
    BASE / "manuscript/09_参考文献.md",
    BASE / "manuscript/附录.md",
]

# Global equation counter
_eq_counter = 0


# ── Fonts & styles ──────────────────────────────────────────────────

def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = Pt(22)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _set_east_asian(style.element, "宋体")


def _set_east_asian(style_elem, cn):
    rPr = style_elem.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        style_elem.append(rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")


def _set_run_font(run, cn="宋体", en="Times New Roman", size=12, bold=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        run._element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)


def _add_text_bold(p, text, cn="宋体", en="Times New Roman", size=12):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            _set_run_font(run, cn, en, size, bold=True)
        else:
            clean = part.replace("\\*", "*")
            if clean:
                run = p.add_run(clean)
                _set_run_font(run, cn, en, size)


# ── Inline variable formatting ────────────────────────────────────────

# Matches inline math $...$ and known variable patterns (longest first)
INLINE_VAR_RE = re.compile(
    r'\$[^$]+\$'                        # $inline math$
    r'|DU_\{?kw[_\\]?ln\}?'            # DU_kw_ln
    r'|DU_\{?sub[_\\]?ln\}?'           # DU_sub_ln
    r'|DU_\{?kw\}?'                    # DU_kw
    r'|DU_\{?llm\}?'                   # DU_llm
    r'|DU_\{?sub\}?'                   # DU_sub
    r'|DU_\{i,t\+1\}'                  # DU_{i,t+1}
    r'|DU_\{i,t\}'                     # DU_{i,t}
    r'|(?<![A-Za-z_])DU(?![A-Za-z_])'  # standalone DU
    r'|data_(?:stock|dev|app|value|gov)' # dimension labels
    r'|PriceDelay'                      # PriceDelay
    r'|SYNCH'                           # SYNCH
    r'|llm_binary'                      # llm_binary
    r'|llm_score'                       # llm_score
    r'|Amihud(?:_w)?'                   # Amihud / Amihud_w
    r'|RetVol(?:_w)?'                   # RetVol / RetVol_w
    r'|Analyst'                         # Analyst
    r'|kw_total'                        # kw_total
    r'|Top1Share'                       # Top1Share
    r'|IndepRatio'                      # IndepRatio
    r'|TobinQ'                          # TobinQ
    r'|(?<![A-Za-z])Size(?![A-Za-z])'  # Size
    r'|(?<![A-Za-z])Lev(?![A-Za-z])'   # Lev
    r'|(?<![A-Za-z])ROA(?![A-Za-z])'   # ROA
    r'|(?<![A-Za-z])Age(?![A-Za-z])'   # Age
    r'|(?<![A-Za-z])Growth(?![A-Za-z])' # Growth
    r'|(?<![A-Za-z])Dual(?![A-Za-z])'  # Dual
    r'|(?<![A-Za-z])SOE(?![A-Za-z])'   # SOE
    r'|(?<![A-Za-z])CFO(?![A-Za-z])'   # CFO
    r'|(?<![A-Za-z])IMR(?![A-Za-z])'   # IMR
    r'|(?<![A-Za-z])Rmax(?![A-Za-z])'  # Rmax
    r'|(?<![A-Za-z])R²'                # R²
)


def _read_latex_group(latex, pos):
    """Read {group} or single char from LaTeX string. Returns (text, chars_consumed)."""
    if pos >= len(latex):
        return ('', 0)
    if latex[pos] == '{':
        end = latex.find('}', pos + 1)
        if end == -1:
            return (latex[pos + 1:], len(latex) - pos)
        return (latex[pos + 1:end], end - pos + 1)
    return (latex[pos], 1)


def _set_italic_sub(run, en, size, subscript=False, superscript=False, bold=False):
    """Set run to italic with optional subscript/superscript."""
    run.font.name = en
    run.font.size = Pt(size)
    run.italic = True
    run.bold = bold
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        run._element.insert(0, rPr)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), "宋体")
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    if subscript:
        run.font.subscript = True
    if superscript:
        run.font.superscript = True


def _add_inline_math_runs(p, latex, en="Times New Roman", size=12):
    """Parse simple inline LaTeX ($..$ content) and add formatted runs."""
    pos = 0
    while pos < len(latex):
        ch = latex[pos]
        if ch == '\\':
            j = pos + 1
            while j < len(latex) and latex[j].isalpha():
                j += 1
            cmd = latex[pos:j]
            pos = j
            if cmd == '\\text':
                text, consumed = _read_latex_group(latex, pos)
                pos += consumed
                run = p.add_run(text)
                _set_run_font(run, "宋体", en, size)
                continue
            char = GREEK_MAP.get(cmd, SYMBOL_MAP.get(cmd, cmd.lstrip('\\')))
            sub_text, sup_text = None, None
            while pos < len(latex) and latex[pos] in ('_', '^'):
                marker = latex[pos]
                pos += 1
                text, consumed = _read_latex_group(latex, pos)
                pos += consumed
                if marker == '_':
                    sub_text = text
                else:
                    sup_text = text
            run = p.add_run(char)
            _set_italic_sub(run, en, size)
            if sub_text:
                rs = p.add_run(sub_text)
                _set_italic_sub(rs, en, size, subscript=True)
            if sup_text:
                rs = p.add_run(sup_text)
                _set_italic_sub(rs, en, size, superscript=True)
        elif ch.isalpha():
            pos += 1
            sub_text, sup_text = None, None
            while pos < len(latex) and latex[pos] in ('_', '^'):
                marker = latex[pos]
                pos += 1
                text, consumed = _read_latex_group(latex, pos)
                pos += consumed
                if marker == '_':
                    sub_text = text
                else:
                    sup_text = text
            run = p.add_run(ch)
            _set_italic_sub(run, en, size)
            if sub_text:
                rs = p.add_run(sub_text)
                _set_italic_sub(rs, en, size, subscript=True)
            if sup_text:
                rs = p.add_run(sup_text)
                _set_italic_sub(rs, en, size, superscript=True)
        elif ch in ' \t':
            pos += 1
        else:
            run = p.add_run(ch)
            _set_run_font(run, "宋体", en, size)
            pos += 1


def _add_variable_runs(p, var_text, en="Times New Roman", size=12, bold=False):
    """Add formatted runs for a matched variable pattern."""
    # $...$  inline math
    if var_text.startswith('$') and var_text.endswith('$'):
        _add_inline_math_runs(p, var_text[1:-1], en, size)
        return

    # DU variants with subscript
    m = re.match(r'^DU_\{?([^}]+)\}?$', var_text)
    if m:
        sub = m.group(1).replace('\\', '')
        run = p.add_run('DU')
        _set_italic_sub(run, en, size, bold=bold)
        rs = p.add_run(sub)
        _set_italic_sub(rs, en, size, subscript=True, bold=bold)
        return

    # Standalone DU
    if var_text.strip() == 'DU':
        run = p.add_run('DU')
        _set_italic_sub(run, en, size, bold=bold)
        return

    # data_xxx dimension labels
    m = re.match(r'^data_(stock|dev|app|value|gov)$', var_text)
    if m:
        run = p.add_run('data')
        _set_italic_sub(run, en, size, bold=bold)
        rs = p.add_run(m.group(1))
        _set_italic_sub(rs, en, size, subscript=True, bold=bold)
        return

    # llm_xxx
    m = re.match(r'^llm_(binary|score)$', var_text)
    if m:
        run = p.add_run('llm')
        _set_italic_sub(run, en, size, bold=bold)
        rs = p.add_run(m.group(1))
        _set_italic_sub(rs, en, size, subscript=True, bold=bold)
        return

    # kw_total
    if var_text == 'kw_total':
        run = p.add_run('kw')
        _set_italic_sub(run, en, size, bold=bold)
        rs = p.add_run('total')
        _set_italic_sub(rs, en, size, subscript=True, bold=bold)
        return

    # Amihud_w / RetVol_w
    m = re.match(r'^(Amihud|RetVol)(_w)?$', var_text)
    if m:
        run = p.add_run(m.group(1))
        _set_italic_sub(run, en, size, bold=bold)
        if m.group(2):
            rs = p.add_run('w')
            _set_italic_sub(rs, en, size, subscript=True, bold=bold)
        return

    # R²
    if var_text.endswith('²'):
        base = var_text[:-1]
        run = p.add_run(base)
        _set_italic_sub(run, en, size, bold=bold)
        rs = p.add_run('2')
        _set_italic_sub(rs, en, size, superscript=True, bold=bold)
        return

    # Rmax
    if var_text == 'Rmax':
        run = p.add_run('R')
        _set_italic_sub(run, en, size, bold=bold)
        rs = p.add_run('max')
        _set_italic_sub(rs, en, size, subscript=True, bold=bold)
        return

    # Simple italic: PriceDelay, SYNCH
    run = p.add_run(var_text)
    _set_italic_sub(run, en, size, bold=bold)


def _add_var_segments(p, text, cn, en, size, bold=False):
    """Split text on variable patterns and add formatted + normal runs."""
    last = 0
    for m in INLINE_VAR_RE.finditer(text):
        if m.start() > last:
            run = p.add_run(text[last:m.start()])
            _set_run_font(run, cn, en, size, bold=bold)
        _add_variable_runs(p, m.group(), en, size, bold=bold)
        last = m.end()
    if last < len(text):
        run = p.add_run(text[last:])
        _set_run_font(run, cn, en, size, bold=bold)


def _add_rich_text(p, text, cn="宋体", en="Times New Roman", size=12):
    """Add paragraph text with bold markers, italic variables, and subscripts."""
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            _add_var_segments(p, part[2:-2], cn, en, size, bold=True)
        else:
            clean = part.replace("\\*", "*")
            if clean:
                _add_var_segments(p, clean, cn, en, size, bold=False)


# ── Paragraph helpers ───────────────────────────────────────────────

def add_heading1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, "黑体", "Times New Roman", 15, bold=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, "黑体", "Times New Roman", 14, bold=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_heading2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, "黑体", "Times New Roman", 14, bold=False)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_heading3(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, "黑体", "Times New Roman", 12, bold=False)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    return p


def add_normal(doc, text, indent=True, size=12, cn="宋体"):
    p = doc.add_paragraph()
    _add_rich_text(p, text, cn, "Times New Roman", size)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = Pt(22)
    return p


def add_note(doc, text):
    return add_normal(doc, text, indent=False, size=10)


# ── OMML Formula Support ───────────────────────────────────────────

GREEK_MAP = {
    "\\alpha": "\u03B1", "\\beta": "\u03B2", "\\gamma": "\u03B3",
    "\\delta": "\u03B4", "\\epsilon": "\u03B5", "\\varepsilon": "\u03B5",
    "\\mu": "\u03BC", "\\lambda": "\u03BB", "\\sigma": "\u03C3",
    "\\pi": "\u03C0", "\\rho": "\u03C1", "\\theta": "\u03B8",
}

SYMBOL_MAP = {
    "\\cdot": "\u00B7", "\\times": "\u00D7", "\\ldots": "\u2026",
    "\\infty": "\u221E", "\\leq": "\u2264", "\\geq": "\u2265",
    "\\neq": "\u2260", "\\approx": "\u2248", "\\sim": "\u223C",
}


def _m_elem(tag, text=None):
    """Create OMML element with optional text child."""
    e = OxmlElement(f"m:{tag}")
    if text is not None:
        t = OxmlElement("m:t")
        t.text = text
        # preserve spaces
        t.set(qn("xml:space"), "preserve")
        e.append(t)
    return e


def _m_run(text, italic=True, bold=False):
    """Create OMML math run <m:r> with text."""
    mr = OxmlElement("m:r")
    # Run properties
    if not italic or bold:
        rPr = OxmlElement("m:rPr")
        if not italic:
            sty = OxmlElement("m:sty")
            sty.set(qn("m:val"), "b" if bold else "p")  # p = plain (not italic)
            rPr.append(sty)
        mr.append(rPr)
    # Word run properties for font
    wrPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Cambria Math")
    rFonts.set(qn("w:hAnsi"), "Cambria Math")
    rFonts.set(qn("w:eastAsia"), "宋体")
    wrPr.append(rFonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "21")  # 10.5pt in half-points
    wrPr.append(sz)
    mr.append(wrPr)
    # Text
    mt = OxmlElement("m:t")
    mt.text = text
    mt.set(qn("xml:space"), "preserve")
    mr.append(mt)
    return mr


def _m_frac(num_elems, den_elems):
    """Create fraction <m:f>."""
    f = OxmlElement("m:f")
    num = OxmlElement("m:num")
    for e in num_elems:
        num.append(e)
    den = OxmlElement("m:den")
    for e in den_elems:
        den.append(e)
    f.append(num)
    f.append(den)
    return f


def _m_sub(base_elems, sub_elems):
    """Create subscript <m:sSub>."""
    s = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    for el in base_elems:
        e.append(el)
    sub = OxmlElement("m:sub")
    for el in sub_elems:
        sub.append(el)
    s.append(e)
    s.append(sub)
    return s


def _m_sup(base_elems, sup_elems):
    """Create superscript <m:sSup>."""
    s = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    for el in base_elems:
        e.append(el)
    sup = OxmlElement("m:sup")
    for el in sup_elems:
        sup.append(el)
    s.append(e)
    s.append(sup)
    return s


def _m_subsup(base_elems, sub_elems, sup_elems):
    """Create sub-superscript <m:sSubSup>."""
    s = OxmlElement("m:sSubSup")
    e = OxmlElement("m:e")
    for el in base_elems:
        e.append(el)
    sub = OxmlElement("m:sub")
    for el in sub_elems:
        sub.append(el)
    sup = OxmlElement("m:sup")
    for el in sup_elems:
        sup.append(el)
    s.append(e)
    s.append(sub)
    s.append(sup)
    return s


def _m_nary(char, sub_elems, sup_elems, body_elems):
    """Create n-ary operator (sum, product, etc.) <m:nary>."""
    nary = OxmlElement("m:nary")
    # Properties
    naryPr = OxmlElement("m:naryPr")
    chrElem = OxmlElement("m:chr")
    chrElem.set(qn("m:val"), char)
    naryPr.append(chrElem)
    nary.append(naryPr)
    # Sub
    sub = OxmlElement("m:sub")
    for el in sub_elems:
        sub.append(el)
    nary.append(sub)
    # Sup
    sup = OxmlElement("m:sup")
    for el in sup_elems:
        sup.append(el)
    nary.append(sup)
    # Body (e)
    e = OxmlElement("m:e")
    for el in body_elems:
        e.append(el)
    nary.append(e)
    return nary


class LaTeXParser:
    """Minimal LaTeX to OMML converter for the formulas in this paper."""

    def __init__(self, latex):
        self.latex = latex.strip()
        self.pos = 0
        self.eq_num = None  # extracted equation number like "(1)"

    def parse(self):
        """Return list of OMML elements."""
        # Extract equation number if present
        m = re.search(r"\\qquad\s*(\(\d+\))\s*$", self.latex)
        if m:
            self.eq_num = m.group(1)
            self.latex = self.latex[:m.start()].strip()

        self.pos = 0
        return self._parse_seq()

    def _peek(self, n=1):
        return self.latex[self.pos:self.pos + n] if self.pos < len(self.latex) else ""

    def _advance(self, n=1):
        self.pos += n

    def _parse_seq(self, stop_at=None):
        """Parse a sequence of math elements."""
        elems = []
        while self.pos < len(self.latex):
            if stop_at and self._peek() == stop_at:
                break
            old_pos = self.pos
            elems.extend(self._parse_atom())
            if self.pos == old_pos:
                self._advance()  # safety: avoid infinite loop
        return elems

    def _parse_group(self):
        """Parse {group} or single token."""
        if self._peek() == "{":
            self._advance()  # skip {
            result = self._parse_seq(stop_at="}")
            if self._peek() == "}":
                self._advance()  # skip }
            return result if result else [_m_run("", italic=False)]
        elif self.pos >= len(self.latex):
            return [_m_run("", italic=False)]
        else:
            # Single character
            result = self._parse_atom()
            return result if result else [_m_run("", italic=False)]

    def _parse_atom(self):
        """Parse one atom (run, command, subscript base, etc.)."""
        ch = self._peek()

        if ch == "\\":
            return self._parse_command()
        elif ch == "{":
            return self._parse_group()
        elif ch == "_" or ch == "^":
            self._advance()  # skip stray sub/sup markers
            group = self._parse_group()
            return group  # orphan sub/sup, just return the group content
        elif ch in "+-=(),;: ":
            self._advance()
            if ch == " ":
                return []  # skip spaces
            return [_m_run(ch, italic=False)]
        elif ch.isdigit() or ch == ".":
            # Collect consecutive digits
            num = ""
            while self.pos < len(self.latex) and (self._peek().isdigit() or self._peek() == "."):
                num += self._peek()
                self._advance()
            elem = _m_run(num, italic=False)
            return self._maybe_scripts([elem])
        elif ch.isalpha():
            self._advance()
            elem = _m_run(ch, italic=True)
            return self._maybe_scripts([elem])
        else:
            self._advance()
            return [_m_run(ch, italic=False)]

    def _parse_command(self):
        """Parse \\command."""
        # Read command name
        j = self.pos + 1
        while j < len(self.latex) and self.latex[j].isalpha():
            j += 1
        cmd = self.latex[self.pos:j]
        self._advance(j - self.pos)

        if cmd == "\\frac":
            num = self._parse_group()
            den = self._parse_group()
            elem = _m_frac(num, den)
            return self._maybe_scripts([elem])
        elif cmd == "\\text":
            content = self._read_brace_text()
            return [_m_run(content, italic=False)]
        elif cmd == "\\sum":
            return self._parse_nary("\u2211")
        elif cmd == "\\prod":
            return self._parse_nary("\u220F")
        elif cmd in GREEK_MAP:
            elem = _m_run(GREEK_MAP[cmd], italic=True)
            return self._maybe_scripts([elem])
        elif cmd in SYMBOL_MAP:
            return [_m_run(SYMBOL_MAP[cmd], italic=False)]
        elif cmd == "\\qquad":
            return []  # handled separately
        else:
            # Unknown command, render as text
            return [_m_run(cmd.lstrip("\\"), italic=True)]

    def _parse_nary(self, char):
        """Parse n-ary like \\sum_{a}^{b} followed by body."""
        sub_elems = []
        sup_elems = []
        if self._peek() == "_":
            self._advance()
            sub_elems = self._parse_group()
        if self._peek() == "^":
            self._advance()
            sup_elems = self._parse_group()
        # Body: parse just the next atom (with potential scripts)
        body = []
        if self.pos < len(self.latex) and self._peek() not in "+=-)}":
            old_pos = self.pos
            body.extend(self._parse_atom())
            if self.pos == old_pos:
                self._advance()
        return [_m_nary(char, sub_elems, sup_elems, body)]

    def _maybe_scripts(self, base_elems):
        """Check for following _ and ^ and apply sub/superscript."""
        sub = None
        sup = None
        while self._peek() in ("_", "^"):
            if self._peek() == "_":
                self._advance()
                sub = self._parse_group()
            elif self._peek() == "^":
                self._advance()
                sup = self._parse_group()
        if sub and sup:
            return [_m_subsup(base_elems, sub, sup)]
        elif sub:
            return [_m_sub(base_elems, sub)]
        elif sup:
            return [_m_sup(base_elems, sup)]
        return base_elems

    def _read_brace_text(self):
        """Read text inside { } without parsing LaTeX."""
        if self._peek() != "{":
            return ""
        self._advance()  # skip {
        depth = 1
        text = ""
        while self.pos < len(self.latex) and depth > 0:
            ch = self._peek()
            if ch == "{":
                depth += 1
            elif ch == "}":
                depth -= 1
                if depth == 0:
                    self._advance()
                    return text
            text += ch
            self._advance()
        return text


def add_omml_formula(doc, latex_str):
    """Add a display equation with OMML and optional equation number."""
    global _eq_counter

    try:
        parser = LaTeXParser(latex_str)
        math_elems = parser.parse()
        if not math_elems:
            math_elems = [_m_run(latex_str, italic=True)]
    except Exception as e:
        print(f"  OMML parse error for: {latex_str[:60]}... -> {e}", flush=True)
        # Fallback: add as plain italic text
        p = doc.add_paragraph()
        run = p.add_run(latex_str)
        _set_run_font(run, "宋体", "Times New Roman", 11)
        run.italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        return p

    # Create paragraph
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    # Set up paragraph properties
    pPr = p._element.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p._element.insert(0, pPr)

    if parser.eq_num:
        # Equation centered via tabs, number right-aligned
        # Use left-aligned paragraph with center + right tab stops
        jc = pPr.find(qn("w:jc"))
        if jc is not None:
            pPr.remove(jc)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "left")
        pPr.append(jc)

        # Tab stops: center at page center, right at right margin
        tabs = OxmlElement("w:tabs")
        tab_c = OxmlElement("w:tab")
        tab_c.set(qn("w:val"), "center")
        tab_c.set(qn("w:pos"), "4253")  # ~7.5cm center of text area
        tabs.append(tab_c)
        tab_r = OxmlElement("w:tab")
        tab_r.set(qn("w:val"), "right")
        tab_r.set(qn("w:pos"), "8505")  # ~15cm right margin
        tabs.append(tab_r)
        pPr.append(tabs)

        # Tab to center -> inline oMath -> Tab to right -> eq number
        tab_run1 = p.add_run("\t")
        _set_run_font(tab_run1, "宋体", "Times New Roman", 11)

        oMath = OxmlElement("m:oMath")
        for elem in math_elems:
            oMath.append(elem)
        p._element.append(oMath)

        tab_run2 = p.add_run("\t")
        _set_run_font(tab_run2, "宋体", "Times New Roman", 11)
        num_run = p.add_run(parser.eq_num)
        _set_run_font(num_run, "宋体", "Times New Roman", 11)
    else:
        # Simple centered equation using oMathPara with explicit centering
        oMathPara = OxmlElement("m:oMathPara")
        oMathParaPr = OxmlElement("m:oMathParaPr")
        jc_m = OxmlElement("m:jc")
        jc_m.set(qn("m:val"), "center")
        oMathParaPr.append(jc_m)
        oMathPara.append(oMathParaPr)
        oMath = OxmlElement("m:oMath")
        for elem in math_elems:
            oMath.append(elem)
        oMathPara.append(oMath)
        p._element.append(oMathPara)

    return p


# ── Three-line table (三线表) ───────────────────────────────────────

def _set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Set cell borders. Each arg is dict with val, sz, color keys."""
    tc = cell._tc
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr)
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is not None:
        tcPr.remove(tcBorders)
    tcBorders = OxmlElement("w:tcBorders")
    for edge_name, props in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if props:
            elem = OxmlElement(f"w:{edge_name}")
            for k, v in props.items():
                elem.set(qn(f"w:{k}"), str(v))
            tcBorders.append(elem)
    tcPr.append(tcBorders)


def _clear_table_borders(table):
    """Remove default table borders."""
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        return
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is not None:
        tblPr.remove(borders)
    # Set empty borders
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "none")
        elem.set(qn("w:sz"), "0")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "auto")
        borders.append(elem)
    tblPr.append(borders)


def parse_md_table(lines):
    rows = []
    for line in lines:
        if '-' in line and re.match(r'^[\s|:\-]+$', line):
            continue
        cells = [c.strip() for c in line.split("|")]
        if cells and cells[0] == "":
            cells = cells[1:]
        if cells and cells[-1] == "":
            cells = cells[:-1]
        if cells:
            rows.append(cells)
    return rows


def _count_header_rows(rows):
    """Auto-detect number of header rows by finding first data row."""
    for i, row in enumerate(rows):
        for cell in row[1:]:
            c = cell.strip()
            if not c:
                continue
            if '***' in c or '**' in c:
                return max(1, i)
            if c in ('YES', 'NO', '否', '是'):
                return max(1, i)
            if re.match(r'^[-+]?0\.\d', c):
                return max(1, i)
            if re.match(r'^\(0\.', c):
                return max(1, i)
    return 1


def add_three_line_table(doc, rows):
    """Add a 三线表 formatted table with auto-detected multi-row header."""
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < ncols:
            r.append("")

    nrows = len(rows)
    hdr_count = _count_header_rows(rows)

    table = doc.add_table(rows=nrows, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _clear_table_borders(table)

    THICK = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    THIN = {"val": "single", "sz": "4", "space": "0", "color": "000000"}
    NONE = {"val": "none", "sz": "0", "space": "0", "color": "auto"}

    # Detect "辅助检验" divider row for endogeneity table
    aux_row = None
    for ri, row in enumerate(rows):
        if row[0].strip() in ('辅助检验',):
            aux_row = ri
            break

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            cell = table.cell(i, j)
            p = cell.paragraphs[0]
            p.clear()

            bold = (i < hdr_count)
            _add_var_segments(p, cell_text, "宋体", "Times New Roman", 10, bold=bold)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = None
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

            # Three-line borders
            top_b = NONE
            bottom_b = NONE

            if i == 0:
                top_b = THICK
            if i == hdr_count - 1:
                bottom_b = THIN
            if i == nrows - 1:
                bottom_b = THICK
            if aux_row and i == aux_row:
                top_b = THIN
                bottom_b = THIN

            _set_cell_borders(cell,
                              top=top_b, bottom=bottom_b,
                              left=NONE, right=NONE)

    return table


# ── Abstract processing ─────────────────────────────────────────────

def process_abstract(doc, fp):
    text = fp.read_text("utf-8")
    lines = text.split("\n")

    title = ""
    subtitle = ""
    abstract = ""
    keywords = ""
    highlights = []
    mode = None

    for line in lines:
        s = line.strip()
        if not s:
            continue
        if s.startswith("# ") and "Highlights" not in s and "摘要" not in s:
            title = s.lstrip("# ").strip()
            mode = None
            continue
        # Subtitle line (starts with ——)
        if s.startswith("\u2014\u2014") or s.startswith("——"):
            subtitle = s
            continue
        if "## Highlights" in s:
            mode = "hl"
            continue
        if "## 摘要" in s:
            mode = "abs"
            continue
        if s.startswith("**关键词**"):
            keywords = s.replace("**关键词**：", "").replace("**关键词**:", "").strip()
            mode = None
            continue
        if mode == "hl" and s.startswith("- "):
            highlights.append(s[2:])
        elif mode == "hl" and s.startswith("#"):
            mode = None
        elif mode == "abs" or (mode is None and not s.startswith("#") and not s.startswith("-")):
            if not title:
                title = s
            else:
                abstract += s

    # Write Highlights
    add_heading1(doc, "Highlights")
    for h in highlights:
        add_normal(doc, "- " + h, indent=False, size=12)

    doc.add_paragraph()
    add_heading1(doc, title)

    # Subtitle
    if subtitle:
        add_subtitle(doc, subtitle)

    # 摘要 label
    p = doc.add_paragraph()
    run = p.add_run("摘要")
    _set_run_font(run, "黑体", "Times New Roman", 12, bold=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_normal(doc, abstract)

    p = doc.add_paragraph()
    run = p.add_run("关键词：")
    _set_run_font(run, "黑体", "Times New Roman", 12, bold=False)
    run2 = p.add_run(keywords)
    _set_run_font(run2, "宋体", "Times New Roman", 12)
    doc.add_paragraph()


# ── Chapter processing ──────────────────────────────────────────────

def process_chapter(doc, fp):
    text = fp.read_text("utf-8")
    lines = text.split("\n")
    n = len(lines)
    i = 0

    while i < n:
        s = lines[i].strip()

        if not s:
            i += 1
            continue

        # ## heading
        if s.startswith("## ") and not s.startswith("### "):
            add_heading2(doc, s.lstrip("# ").strip())
            i += 1
            continue

        # ### heading
        if s.startswith("### "):
            add_heading3(doc, s.lstrip("# ").strip())
            i += 1
            continue

        # **bold heading** (standalone)
        if s.startswith("**") and s.endswith("**") and len(s) < 100:
            inner = s[2:-2]
            if re.match(r'^[表图]\d', inner) or inner in (
                    '数据要素利用关键词体系',):
                # Table/figure title - centered, songti
                p = doc.add_paragraph()
                _add_rich_text(p, inner, "宋体", "Times New Roman", 10.5)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = Pt(22)
            else:
                p = doc.add_paragraph()
                run = p.add_run(inner)
                _set_run_font(run, "黑体", "Times New Roman", 12, bold=False)
            i += 1
            continue

        # Table: line with | and next line has ---
        if "|" in s:
            next_s = lines[i + 1].strip() if i + 1 < n else ""
            if "---" in next_s or "---" in s:
                tbl_lines = []
                while i < n and "|" in lines[i]:
                    tbl_lines.append(lines[i])
                    i += 1
                rows = parse_md_table(tbl_lines)
                if rows:
                    add_three_line_table(doc, rows)
                continue

        # 注：
        if s.startswith("注：") or s.startswith("注:"):
            add_note(doc, s)
            i += 1
            continue

        # ![alt](path) image
        img_m = re.match(r'^!\[.*?\]\((.+?)\)$', s)
        if img_m:
            img_path = BASE / img_m.group(1)
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(img_path), width=Cm(14))
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
            else:
                print(f"  Image not found: {img_path}", flush=True)
            i += 1
            continue

        # $$ formula (display equation)
        if s.startswith("$$"):
            formula = s.replace("$$", "").strip()
            if not formula:
                # Multi-line formula
                i += 1
                parts = []
                while i < n and not lines[i].strip().startswith("$$"):
                    parts.append(lines[i].strip())
                    i += 1
                formula = " ".join(parts)
                if i < n:
                    i += 1  # skip closing $$
            else:
                # Single-line: $$formula$$
                formula = formula.rstrip("$").strip()
                i += 1

            if formula:
                add_omml_formula(doc, formula)
            continue

        # Normal paragraph (inline $var$ handled by _add_rich_text)
        add_normal(doc, s)
        i += 1


# ── References processing ───────────────────────────────────────────

def process_references(doc, fp):
    text = fp.read_text("utf-8")
    add_heading2(doc, "参考文献")
    for line in text.split("\n"):
        s = line.strip()
        if not s or s.startswith("#"):
            continue
        m = re.match(r"^(\d+)\.\s+(.+)$", s)
        if m:
            p = doc.add_paragraph()
            run = p.add_run(f"[{m.group(1)}] {m.group(2)}")
            _set_run_font(run, "宋体", "Times New Roman", 10.5)
            p.paragraph_format.first_line_indent = None
        else:
            add_normal(doc, s, indent=False, size=10.5)


# ── Appendix processing ────────────────────────────────────────────

def process_appendix(doc, fp):
    text = fp.read_text("utf-8")
    lines = text.split("\n")
    n = len(lines)
    i = 0
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("## "):
            add_heading2(doc, s.lstrip("# ").strip())
            i += 1
            continue
        if "|" in s:
            next_s = lines[i + 1].strip() if i + 1 < n else ""
            if "---" in next_s:
                tbl = []
                while i < n and "|" in lines[i]:
                    tbl.append(lines[i])
                    i += 1
                rows = parse_md_table(tbl)
                if rows:
                    add_three_line_table(doc, rows)
                continue
        if s.startswith("注：") or s.startswith("注:"):
            add_note(doc, s)
            i += 1
            continue
        add_normal(doc, s, indent=False)
        i += 1


# ── Main ────────────────────────────────────────────────────────────

def main():
    print("Creating v11 document...", flush=True)
    doc = Document()
    setup_styles(doc)

    if doc.paragraphs:
        doc.paragraphs[0]._element.getparent().remove(doc.paragraphs[0]._element)

    print("Processing 00_摘要...", flush=True)
    process_abstract(doc, MD_FILES[0])

    for fp in MD_FILES[1:8]:
        print(f"Processing {fp.stem}...", flush=True)
        process_chapter(doc, fp)

    print("Processing 09_参考文献...", flush=True)
    process_references(doc, MD_FILES[8])

    print("Processing 附录...", flush=True)
    process_appendix(doc, MD_FILES[9])

    print(f"Saving {OUTPUT.name}...", flush=True)
    doc.save(str(OUTPUT))

    d2 = Document(str(OUTPUT))
    print(f"Done: {len(d2.paragraphs)} paragraphs, {len(d2.tables)} tables")


if __name__ == "__main__":
    main()
